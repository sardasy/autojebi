"""Bounded, lossless table context; coordinates are zero-based, pages one-based.

Adapters do not discover headers or units: callers supply verified header bands
and source-backed units. XLSX accepts an already opened, non-read-only worksheet;
archive/decompression limits belong to its loader. HWPX accepts one XML part,
not a ZIP. No adapter performs OCR, formula evaluation, or approval.
"""

from __future__ import annotations

import json
from collections.abc import Sequence
from math import isfinite
from typing import Annotated, Literal
from xml.etree.ElementTree import Element, ParseError

from defusedxml.common import DefusedXmlException
from defusedxml.ElementTree import fromstring
from openpyxl.utils.cell import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from pydantic import BaseModel, ConfigDict, Field, ValidationError, model_validator

Index = Annotated[int, Field(strict=True, ge=0, le=20_000)]
Text = Annotated[str, Field(max_length=200_000)]
HeaderBand = Annotated[tuple[Index, ...], Field(min_length=1, max_length=32)]


class TableStructureError(ValueError):
    """Explicit rejection, never a truncated successful extraction."""


class _Model(BaseModel):
    model_config = ConfigDict(extra="forbid", frozen=True)


class TableLimits(_Model):
    max_rows: int = Field(default=2_000, ge=1, le=20_000)
    max_cols: int = Field(default=128, ge=1, le=512)
    max_cells: int = Field(default=20_000, ge=1, le=100_000)
    max_text_chars: int = Field(default=200_000, ge=1, le=1_000_000)
    max_xml_bytes: int = Field(default=2_000_000, ge=1, le=10_000_000)
    max_xml_nodes: int = Field(default=50_000, ge=1, le=200_000)
    max_chunk_chars: int = Field(default=32_000, ge=1, le=1_000_000)
    max_chunks: int = Field(default=2_000, ge=1, le=20_000)
    max_output_chars: int = Field(default=2_000_000, ge=1, le=10_000_000)


class SourceLocator(_Model):
    locator: str = Field(min_length=1, max_length=2_000)
    page: int | None = Field(default=None, strict=True, ge=1)
    bbox: tuple[float, float, float, float] | None = None

    @model_validator(mode="after")
    def valid_bbox(self) -> SourceLocator:
        if self.bbox is not None:
            x0, y0, x1, y1 = self.bbox
            if not all(isfinite(v) for v in self.bbox) or x0 > x1 or y0 > y1:
                raise ValueError("invalid_bbox")
        return self


class SourceText(_Model):
    raw_text: Text
    source: SourceLocator


class HeaderPath(_Model):
    col: Index
    headers: tuple[SourceText, ...] = Field(default=(), max_length=32)


class TableCell(_Model):
    row: Index
    col: Index
    row_span: int = Field(default=1, strict=True, ge=1, le=20_000)
    col_span: int = Field(default=1, strict=True, ge=1, le=512)
    raw_text: Text
    source: SourceLocator
    header_paths: tuple[HeaderPath, ...] = Field(default=(), max_length=512)


class ColumnUnit(_Model):
    col: Index
    value: SourceText


class TableIssue(_Model):
    code: str = Field(min_length=1, max_length=100)
    sources: tuple[SourceLocator, ...] = Field(default=(), max_length=100_000)


class TableStructure(_Model):
    format: Literal["xlsx", "hwpx", "pdf"]
    source: SourceLocator
    title: SourceText | None = None
    n_rows: int = Field(strict=True, ge=1, le=20_000)
    n_cols: int = Field(strict=True, ge=1, le=512)
    cells: tuple[TableCell, ...] = Field(max_length=100_000)
    # Each band is one complete header, including repeated page headers.
    header_bands: tuple[HeaderBand, ...] = Field(default=(), max_length=2_000)
    row_label_col: Index | None = 0
    units: tuple[ColumnUnit, ...] = Field(default=(), max_length=512)
    headers_confirmed: bool = False
    issues: tuple[TableIssue, ...] = Field(default=(), max_length=100_000)
    needs_review: bool = False


class TableChunk(_Model):
    row: Index
    kind: Literal["header", "data"]
    text: str
    source_locators: tuple[SourceLocator, ...]
    needs_review: bool


class SpanContext(_Model):
    """Supplemental retrieval context, never a source quote or an approval."""

    status: Literal["matched", "no_match", "ambiguous", "needs_review"]
    text: str = ""
    source_locators: tuple[SourceLocator, ...] = ()
    needs_review: bool = False
    codes: tuple[str, ...] = ()


def _dimensions(rows: int, cols: int, limits: TableLimits) -> None:
    if not (1 <= rows <= limits.max_rows and 1 <= cols <= limits.max_cols):
        raise TableStructureError("table_dimension_limit")
    if rows * cols > limits.max_cells:
        raise TableStructureError("table_cell_limit")


def _source_text(cell: TableCell) -> SourceText:
    return SourceText(raw_text=cell.raw_text, source=cell.source)


def preserve_table_structure(
    table: TableStructure, *, limits: TableLimits | None = None
) -> TableStructure:
    """Validate bounds, retain conflicts, and derive only explicit header paths.

    Safe to call again: caller-provided/previous header paths are recomputed.
    A conflicting grid disables all derived header paths, retaining raw cells.
    """
    limits = limits or TableLimits()
    _dimensions(table.n_rows, table.n_cols, limits)
    if len(table.cells) > limits.max_cells:
        raise TableStructureError("table_cell_limit")
    text_size = sum(len(cell.raw_text) for cell in table.cells)
    text_size += len(table.title.raw_text) if table.title else 0
    text_size += sum(len(unit.value.raw_text) for unit in table.units)
    if text_size > limits.max_text_chars:
        raise TableStructureError("table_text_limit")
    if table.row_label_col is not None and table.row_label_col >= table.n_cols:
        raise TableStructureError("row_label_out_of_bounds")
    if any(unit.col >= table.n_cols for unit in table.units):
        raise TableStructureError("unit_out_of_bounds")
    if len({unit.col for unit in table.units}) != len(table.units):
        raise TableStructureError("conflicting_units")
    rows = [row for band in table.header_bands for row in band]
    if (
        any(not band or tuple(sorted(set(band))) != band for band in table.header_bands)
        or rows != sorted(set(rows))
        or any(row >= table.n_rows for row in rows)
        or any(band != tuple(range(band[0], band[-1] + 1)) for band in table.header_bands)
    ):
        raise TableStructureError("invalid_header_bands")

    issues = list(table.issues)
    issue_keys = set(issues)

    def issue(code: str, *sources: SourceLocator) -> None:
        item = TableIssue(code=code, sources=sources)
        if item not in issue_keys:
            issues.append(item)
            issue_keys.add(item)

    grid: dict[tuple[int, int], TableCell] = {}
    conflict = False
    occupied = 0
    for cell in table.cells:
        if cell.row + cell.row_span > table.n_rows or cell.col + cell.col_span > table.n_cols:
            raise TableStructureError("span_out_of_bounds")
        occupied += cell.row_span * cell.col_span
        if occupied > limits.max_cells:
            raise TableStructureError("table_span_limit")
        overlap = False
        for row in range(cell.row, cell.row + cell.row_span):
            for col in range(cell.col, cell.col + cell.col_span):
                key = row, col
                if key in grid:
                    conflict = overlap = True
                else:
                    grid[key] = cell
        if overlap:
            issue("conflicting_spans", cell.source)
    if len(grid) != table.n_rows * table.n_cols:
        issue("missing_cells", table.source)
    if not table.header_bands or not table.headers_confirmed:
        issue("ambiguous_headers", table.source)

    header_rows = set(rows)
    for row in rows:
        for col in range(table.n_cols):
            header = grid.get((row, col))
            if header is None or not header.raw_text.strip():
                issue("ambiguous_headers", header.source if header else table.source)
    band_rows = [set(band) for band in table.header_bands]
    for cell in table.cells:
        covered = set(range(cell.row, cell.row + cell.row_span))
        if covered & header_rows and not covered <= header_rows:
            conflict = True
            issue("header_data_span", cell.source)
        elif covered & header_rows and not any(covered <= band for band in band_rows):
            conflict = True
            issue("header_band_span", cell.source)

    # No nearest-neighbor inference and no filling from previous nonempty cells.
    resolved = []
    path_cache: dict[tuple[int, int], HeaderPath] = {}
    for cell in sorted(table.cells, key=lambda c: (c.row, c.col)):
        paths = []
        bands = [band for band in table.header_bands if band[-1] < cell.row]
        if bands and cell.row not in header_rows and table.headers_confirmed and not conflict:
            for col in range(cell.col, cell.col + cell.col_span):
                cache_key = bands[-1][-1], col
                if cache_key in path_cache:
                    paths.append(path_cache[cache_key])
                    continue
                headers: list[SourceText] = []
                anchors: set[tuple[int, int]] = set()
                for row in bands[-1]:
                    header = grid.get((row, col))
                    if header is None or not header.raw_text.strip():
                        issue("ambiguous_headers", header.source if header else table.source)
                        headers = []
                        break
                    anchor = header.row, header.col
                    if anchor not in anchors:
                        headers.append(_source_text(header))
                        anchors.add(anchor)
                path = HeaderPath(col=col, headers=tuple(headers))
                path_cache[cache_key] = path
                paths.append(path)
            if len({tuple(h.raw_text for h in p.headers) for p in paths}) > 1:
                issue("merged_value_multiple_headers", cell.source)
        elif cell.row not in header_rows:
            issue("ambiguous_headers", table.source)
        resolved.append(cell.model_copy(update={"header_paths": tuple(paths)}))
    return table.model_copy(
        update={
            "cells": tuple(resolved),
            "issues": tuple(issues),
            "needs_review": table.needs_review or bool(issues),
        }
    )


def table_from_xlsx(
    worksheet: Worksheet,
    *,
    locator: str,
    title: SourceText | None = None,
    header_bands: tuple[tuple[int, ...], ...] = (),
    row_label_col: int | None = 0,
    units: tuple[ColumnUnit, ...] = (),
    limits: TableLimits | None = None,
) -> TableStructure:
    """Preserve a worksheet's merge anchors and every unmerged empty cell.

    Header bands use absolute zero-based worksheet rows, not Excel row numbers.
    Supply only a worksheet whose archive was checked before loading.
    """
    limits = limits or TableLimits()
    if not isinstance(worksheet, Worksheet):
        raise TableStructureError("xlsx_requires_merge_aware_worksheet")
    rows, cols = worksheet.max_row, worksheet.max_column
    _dimensions(rows, cols, limits)
    base = f"{locator}#sheet={worksheet.title}"
    merged = sorted(worksheet.merged_cells.ranges, key=lambda r: (r.min_row, r.min_col, str(r)))
    if len(merged) > limits.max_cells:
        raise TableStructureError("table_cell_limit")
    cells: list[TableCell] = []
    covered: set[tuple[int, int]] = set()
    issues: list[TableIssue] = []
    total_text = 0

    def append(row: int, col: int, row_span: int = 1, col_span: int = 1) -> None:
        nonlocal total_text
        original = worksheet.cell(row, col)
        text = "" if original.value is None else str(original.value)
        total_text += len(text)
        if total_text > limits.max_text_chars:
            raise TableStructureError("table_text_limit")
        address = f"{get_column_letter(col)}{row}"
        if row_span > 1 or col_span > 1:
            address += f":{get_column_letter(col + col_span - 1)}{row + row_span - 1}"
        source = SourceLocator(locator=f"{base}!{address}")
        cells.append(
            TableCell(
                row=row - 1,
                col=col - 1,
                row_span=row_span,
                col_span=col_span,
                raw_text=text,
                source=source,
            )
        )
        if original.data_type in {"f", "e"}:
            issues.append(TableIssue(code="formula_or_error_review", sources=(source,)))

    span_size = 0
    for region in merged:
        height, width = region.max_row - region.min_row + 1, region.max_col - region.min_col + 1
        if region.max_row > rows or region.max_col > cols:
            raise TableStructureError("span_out_of_bounds")
        span_size += height * width
        if span_size > limits.max_cells:
            raise TableStructureError("table_span_limit")
        append(region.min_row, region.min_col, height, width)
        covered.update(
            (r, c)
            for r in range(region.min_row, region.max_row + 1)
            for c in range(region.min_col, region.max_col + 1)
        )
    for row in range(1, rows + 1):
        for col in range(1, cols + 1):
            if (row, col) not in covered:
                append(row, col)
    return preserve_table_structure(
        TableStructure(
            format="xlsx",
            source=SourceLocator(locator=base),
            title=title,
            n_rows=rows,
            n_cols=cols,
            cells=tuple(cells),
            header_bands=header_bands,
            headers_confirmed=bool(header_bands),
            row_label_col=row_label_col,
            units=units,
            issues=tuple(issues),
        ),
        limits=limits,
    )


def _local(element: Element) -> str:
    return element.tag.rsplit("}", 1)[-1]


def _child(element: Element, name: str) -> Element:
    children = [child for child in element if _local(child) == name]
    if len(children) != 1:
        raise TableStructureError("missing_or_conflicting_hwpx_geometry")
    return children[0]


def _integer(element: Element, name: str) -> int:
    try:
        return int(element.attrib[name])
    except (KeyError, ValueError):
        raise TableStructureError("invalid_hwpx_geometry") from None


def _paragraph_text(paragraph: Element) -> str:
    # Only text inside hp:t is content; indentation around run elements is not.
    # Walk iteratively so deeply nested untrusted XML cannot exhaust recursion.
    parts = []
    stack: list[tuple[Element | str, bool]] = [(paragraph, False)]
    while stack:
        element, in_text = stack.pop()
        if isinstance(element, str):
            parts.append(element)
            continue
        name = _local(element)
        in_text = in_text or name == "t"
        if name in {"tab", "lineBreak"}:
            parts.append("\t" if name == "tab" else "\n")
        if in_text and element.text:
            parts.append(element.text)
        for child in reversed(element):
            if in_text and child.tail:
                stack.append((child.tail, True))
            stack.append((child, in_text))
    return "".join(parts)


def table_from_hwpx_xml(
    xml: bytes | str,
    *,
    locator: str,
    table_index: int = 0,
    title: SourceText | None = None,
    header_bands: tuple[tuple[int, ...], ...] = (),
    row_label_col: int | None = 0,
    units: tuple[ColumnUnit, ...] = (),
    page: int | None = None,
    limits: TableLimits | None = None,
) -> TableStructure:
    """Read explicit HWPX cellAddr/cellSpan; page is caller supplied, not guessed."""
    limits = limits or TableLimits()
    if len(xml) > limits.max_xml_bytes:
        raise TableStructureError("xml_byte_limit")
    content = xml.encode("utf-8") if isinstance(xml, str) else xml
    if len(content) > limits.max_xml_bytes:
        raise TableStructureError("xml_byte_limit")
    try:
        root = fromstring(content, forbid_dtd=True, forbid_entities=True, forbid_external=True)
    except (DefusedXmlException, ParseError):
        raise TableStructureError("unsafe_or_invalid_hwpx_xml") from None
    tables = []
    for count, element in enumerate(root.iter(), 1):
        if count > limits.max_xml_nodes:
            raise TableStructureError("xml_node_limit")
        if _local(element) == "tbl":
            tables.append(element)
    if type(table_index) is not int or not 0 <= table_index < len(tables):
        raise TableStructureError("hwpx_table_not_found")
    table = tables[table_index]
    if any(_local(el) == "tbl" for el in table.iter() if el is not table):
        raise TableStructureError("nested_hwpx_table_requires_review")
    rows, cols = _integer(table, "rowCnt"), _integer(table, "colCnt")
    _dimensions(rows, cols, limits)
    base = f"{locator}#table={table_index}"
    cells = []
    total_text = 0
    for tr in table:
        if _local(tr) != "tr":
            continue
        for tc in tr:
            if _local(tc) != "tc":
                continue
            if len(cells) >= limits.max_cells:
                raise TableStructureError("table_cell_limit")
            address, span = _child(tc, "cellAddr"), _child(tc, "cellSpan")
            row, col = _integer(address, "rowAddr"), _integer(address, "colAddr")
            height, width = _integer(span, "rowSpan"), _integer(span, "colSpan")
            if row < 0 or col < 0 or height < 1 or width < 1:
                raise TableStructureError("invalid_hwpx_geometry")
            if row + height > rows or col + width > cols:
                raise TableStructureError("span_out_of_bounds")
            paragraphs = []
            for paragraph in tc.iter():
                if _local(paragraph) == "p":
                    paragraphs.append(_paragraph_text(paragraph))
            text = "\n".join(paragraphs)
            total_text += len(text)
            if total_text > limits.max_text_chars:
                raise TableStructureError("table_text_limit")
            cells.append(
                TableCell(
                    row=row,
                    col=col,
                    row_span=height,
                    col_span=width,
                    raw_text=text,
                    source=SourceLocator(
                        locator=f"{base}/cell={len(cells)};r={row};c={col}", page=page
                    ),
                )
            )
    return preserve_table_structure(
        TableStructure(
            format="hwpx",
            source=SourceLocator(locator=base, page=page),
            title=title,
            n_rows=rows,
            n_cols=cols,
            cells=tuple(cells),
            header_bands=header_bands,
            headers_confirmed=bool(header_bands),
            row_label_col=row_label_col,
            units=units,
        ),
        limits=limits,
    )


def table_from_pdf_cells(
    cells: Sequence[TableCell],
    *,
    locator: str,
    n_rows: int,
    n_cols: int,
    title: SourceText | None = None,
    header_bands: tuple[tuple[int, ...], ...] = (),
    headers_confirmed: bool = False,
    ocr_ambiguous: bool = False,
    row_label_col: int | None = 0,
    units: tuple[ColumnUnit, ...] = (),
    limits: TableLimits | None = None,
) -> TableStructure:
    """Adapt caller-extracted PDF/OCR cells without inferring spans or filling blanks."""
    limits = limits or TableLimits()
    _dimensions(n_rows, n_cols, limits)
    if len(cells) > limits.max_cells:
        raise TableStructureError("table_cell_limit")
    issues = (TableIssue(code="ambiguous_ocr"),) if ocr_ambiguous else ()
    if any(cell.source.page is None or cell.source.bbox is None for cell in cells):
        issues += (TableIssue(code="missing_pdf_geometry"),)
    return preserve_table_structure(
        TableStructure(
            format="pdf",
            source=SourceLocator(locator=locator),
            title=title,
            n_rows=n_rows,
            n_cols=n_cols,
            cells=tuple(cells),
            header_bands=header_bands,
            headers_confirmed=headers_confirmed and not ocr_ambiguous,
            row_label_col=row_label_col,
            units=units,
            issues=issues,
        ),
        limits=limits,
    )


def build_table_chunks(
    table: TableStructure, *, limits: TableLimits | None = None
) -> tuple[TableChunk, ...]:
    """Return atomic JSON row chunks with source-backed context, or raise.

    Every anchor (including empty cells and repeated headers) occurs once.
    Covered merge positions are not invented cells. Oversized atomic rows fail
    explicitly; callers must not treat a TableStructureError as successful output.
    Row labels may reference an explicit vertical merge, never a preceding blank.
    Chunk text is contextual serialization, not a contiguous verbatim quote.
    """
    limits = limits or TableLimits()
    table = preserve_table_structure(table, limits=limits)
    grouped: dict[int, list[TableCell]] = {}
    row_labels: dict[int, list[TableCell]] = {}
    for cell in table.cells:
        grouped.setdefault(cell.row, []).append(cell)
        if cell.col == table.row_label_col:
            for row in range(cell.row, cell.row + cell.row_span):
                row_labels.setdefault(row, []).append(cell)
    if len(grouped) > limits.max_chunks:
        raise TableStructureError("table_chunk_count_limit")
    chunks = []
    total_chars = 0
    header_rows = {row for band in table.header_bands for row in band}
    for row, cells in sorted(grouped.items()):
        labels = row_labels.get(row, [])
        label = _source_text(labels[0]) if len(labels) == 1 else None
        kind = "header" if row in header_rows else "data"
        payload = {
            "table_source": table.source.model_dump(mode="json"),
            "title": table.title.model_dump(mode="json") if table.title else None,
            "row": row,
            "kind": kind,
            "row_label": label.model_dump(mode="json") if label else None,
            "units": [unit.model_dump(mode="json") for unit in table.units],
            "cells": [cell.model_dump(mode="json") for cell in cells],
            "needs_review": table.needs_review,
            "issues": [issue.model_dump(mode="json") for issue in table.issues],
        }
        parts = []
        size = 0
        encoder = json.JSONEncoder(ensure_ascii=False, separators=(",", ":"), allow_nan=False)
        for part in encoder.iterencode(payload):
            size += len(part)
            if size > limits.max_chunk_chars:
                raise TableStructureError("table_chunk_size_limit")
            if total_chars + size > limits.max_output_chars:
                raise TableStructureError("table_chunk_output_limit")
            parts.append(part)
        text = "".join(parts)
        total_chars += size
        sources = [table.source]
        if table.title:
            sources.append(table.title.source)
        if label:
            sources.append(label.source)
        sources.extend(unit.value.source for unit in table.units)
        for cell in cells:
            sources.append(cell.source)
            sources.extend(header.source for path in cell.header_paths for header in path.headers)
        sources.extend(source for issue in table.issues for source in issue.sources)
        chunks.append(
            TableChunk(
                row=row,
                kind=kind,
                text=text,
                source_locators=tuple(dict.fromkeys(sources)),
                needs_review=table.needs_review,
            )
        )
    return tuple(chunks)


def _check_context_metadata(value: object, limits: TableLimits) -> None:
    """Bound JSON input before Pydantic constructs potentially nested objects."""
    remaining_nodes = min(200_000, limits.max_cells * 64)
    remaining_chars = limits.max_output_chars
    stack = [(value, 0)]
    while stack:
        item, depth = stack.pop()
        remaining_nodes -= 1
        if remaining_nodes < 0 or depth > 32:
            raise TableStructureError("span_context_metadata_limit")
        if isinstance(item, str):
            remaining_chars -= len(item)
        elif type(item) in (dict, list, tuple):
            if len(item) + len(stack) > remaining_nodes:
                raise TableStructureError("span_context_metadata_limit")
            if isinstance(item, dict):
                if any(not isinstance(key, str) for key in item):
                    raise TableStructureError("invalid_table_metadata")
                remaining_chars -= sum(len(key) for key in item)
                stack.extend((child, depth + 1) for child in item.values())
            else:
                stack.extend((child, depth + 1) for child in item)
        elif item is not None and type(item) not in (bool, int, float):
            raise TableStructureError("invalid_table_metadata")
        if remaining_chars < 0:
            raise TableStructureError("span_context_metadata_limit")


def _context_token_char(value: str) -> bool:
    # Lexical boundary only: never parse/normalize/compare numbers or units.
    return value.isalnum() or value in "_.,+-/%"


def context_for_span(
    extraction: dict, text: str, start: int, end: int, *, limits: TableLimits | None = None
) -> dict:
    """Return supplemental JSON context for uniquely anchored, complete cells.

    A raw value must belong to exactly one cell across ALL supplied tables and
    occur exactly once in the full text, wholly inside [start, end). Overlapping
    candidate matches, partial/unknown extraction, or unconfirmed table geometry
    abstain. Literal boundaries prevent matching "4" inside "400"; no numeric
    interpretation, whitespace normalization, nearest-cell mapping, or offsets
    are emitted. Stored header paths are recomputed from explicit spans.

    Output keys: status, text, source_locators, needs_review, codes. Nonempty text
    is JSON containing source-backed cell/title/row_label/header_paths/units.
    "matched" means structural anchoring only, NOT reviewed/approved/valid spec.
    Any invalid input or exceeded bound returns empty, review-required context.
    Input text/tables are never mutated; no partially built context is returned.
    """
    limits = limits or TableLimits()

    def empty(status: str, code: str, review: bool = True) -> dict:
        return SpanContext(status=status, needs_review=review, codes=(code,)).model_dump(
            mode="json"
        )

    if (
        not isinstance(extraction, dict)
        or not isinstance(text, str)
        or type(start) is not int
        or type(end) is not int
        or not 0 <= start < end <= len(text)
    ):
        return empty("needs_review", "invalid_source_span")
    if len(text) > limits.max_text_chars:
        return empty("needs_review", "span_context_text_limit")
    raw_tables = extraction.get("tables", [])
    summary = extraction.get("table_structure", {})
    pages = extraction.get("pages", [])
    try:
        if not isinstance(raw_tables, list) or len(raw_tables) > 100:
            raise TableStructureError("span_context_table_limit")
        raw_cell_count = 0
        for raw_table in raw_tables:
            if not isinstance(raw_table, dict) or not isinstance(
                raw_table.get("cells"), list | tuple
            ):
                raise TableStructureError("invalid_table_metadata")
            raw_cell_count += len(raw_table["cells"])
            if raw_cell_count > limits.max_cells:
                raise TableStructureError("span_context_cell_limit")
        _check_context_metadata([raw_tables, summary, pages], limits)
        if not isinstance(summary, dict) or not isinstance(pages, list):
            raise TableStructureError("invalid_table_metadata")
        if (
            summary.get("version") != 1
            or summary.get("status") not in {"preserved", "needs_review", "not_detected"}
            or summary.get("issues")
            or summary.get("omitted_issue_count")
        ):
            return empty("needs_review", "incomplete_table_extraction")
        if (
            summary.get("detected_count") != len(raw_tables)
            or summary.get("preserved_count") != len(raw_tables)
            or raw_tables
            and summary.get("status") == "not_detected"
        ):
            return empty("needs_review", "incomplete_table_extraction")
        if any(not isinstance(page, dict) for page in pages):
            raise TableStructureError("invalid_table_metadata")
        # Without verified page-to-text spans, an OCR page cannot be excluded.
        if any(page.get("method") == "ocr" or page.get("review_required") for page in pages):
            return empty("needs_review", "ocr_review_required")
        tables: list[TableStructure] = []
        values: dict[str, list[tuple[int, TableCell]]] = {}
        sources: set[str] = set()
        count, chars = 0, 0
        for index, raw_table in enumerate(raw_tables):
            table = TableStructure.model_validate(raw_table)
            count += len(table.cells)
            chars += sum(len(cell.raw_text) for cell in table.cells)
            if count > limits.max_cells or chars > limits.max_text_chars:
                raise TableStructureError("span_context_cell_limit")
            table = preserve_table_structure(table, limits=limits)
            tables.append(table)
            for cell in table.cells:
                if cell.source.locator in sources:
                    return empty("ambiguous", "duplicate_cell_locator")
                sources.add(cell.source.locator)
                if cell.raw_text.strip():
                    values.setdefault(cell.raw_text, []).append((index, cell))

        span = text[start:end]
        work = limits.max_text_chars * 100
        matches: list[tuple[int, int, int, TableCell]] = []
        for raw, candidates in values.items():
            work -= len(span) + len(raw)
            if work < 0:
                raise TableStructureError("span_context_scan_limit")
            if raw not in span:
                continue
            work -= len(text) * 2 + len(raw)
            if work < 0:
                raise TableStructureError("span_context_scan_limit")
            position = text.find(raw)
            stop = position + len(raw)
            if (
                position and _context_token_char(raw[0]) and _context_token_char(text[position - 1])
            ) or (
                stop < len(text)
                and _context_token_char(raw[-1])
                and _context_token_char(text[stop])
            ):
                continue
            if not start <= position or position + len(raw) > end:
                # A later occurrence inside the span still has ambiguous origin.
                return empty("ambiguous", "repeated_text_match")
            if text.find(raw, position + 1) != -1:
                return empty("ambiguous", "repeated_text_match")
            if len(candidates) != 1:
                return empty("ambiguous", "duplicate_cell_value")
            table_index, cell = candidates[0]
            if any(cell.row in band for band in tables[table_index].header_bands):
                continue
            matches.append((position, stop, table_index, cell))
            if len(matches) > 32:
                raise TableStructureError("span_context_match_limit")
        if not matches:
            return empty("no_match", "no_unique_cell_anchor", False)
        matches.sort(key=lambda match: match[0])
        if any(left[1] > right[0] for left, right in zip(matches, matches[1:], strict=False)):
            return empty("ambiguous", "overlapping_cell_matches")
        contexts, locators = [], []
        for _, _, table_index, cell in matches:
            table = tables[table_index]
            if (
                table.needs_review
                or not table.headers_confirmed
                or not cell.header_paths
                or any(not path.headers for path in cell.header_paths)
            ):
                return empty("needs_review", "table_structure_requires_review")
            labels = [
                candidate
                for candidate in table.cells
                if candidate.col == table.row_label_col
                and candidate.row <= cell.row < candidate.row + candidate.row_span
            ]
            label = _source_text(labels[0]) if len(labels) == 1 else None
            units = [
                unit for unit in table.units if cell.col <= unit.col < cell.col + cell.col_span
            ]
            contexts.append(
                {
                    "table_source": table.source.model_dump(mode="json"),
                    "cell": _source_text(cell).model_dump(mode="json"),
                    "title": table.title.model_dump(mode="json") if table.title else None,
                    "row_label": label.model_dump(mode="json") if label else None,
                    "header_paths": [path.model_dump(mode="json") for path in cell.header_paths],
                    "units": [unit.model_dump(mode="json") for unit in units],
                }
            )
            locators.extend((table.source, cell.source))
            if table.title:
                locators.append(table.title.source)
            if label:
                locators.append(label.source)
            locators.extend(header.source for path in cell.header_paths for header in path.headers)
            locators.extend(unit.value.source for unit in units)
        parts, size = [], 0
        for part in json.JSONEncoder(
            ensure_ascii=False, separators=(",", ":"), allow_nan=False
        ).iterencode(contexts):
            size += len(part)
            if size > limits.max_chunk_chars:
                raise TableStructureError("span_context_output_limit")
            parts.append(part)
        review = (
            bool(summary.get("needs_review"))
            or summary.get("status") != "preserved"
            or summary.get("structure_trustworthy") is False
        )
        return SpanContext(
            status="needs_review" if review else "matched",
            text="".join(parts),
            source_locators=tuple(dict.fromkeys(locators)),
            needs_review=review,
            codes=("extraction_requires_review",) if review else (),
        ).model_dump(mode="json")
    except TableStructureError as exc:
        return empty("needs_review", str(exc))
    except (ValidationError, TypeError, ValueError):
        return empty("needs_review", "invalid_table_metadata")
