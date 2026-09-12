"""Bounded spreadsheet/archive extraction, without Office automation or unpacking."""
from __future__ import annotations

import hashlib
import io
import json
import stat
import zipfile
from datetime import date, datetime
from pathlib import PurePosixPath

from defusedxml.ElementTree import fromstring
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.cell_range import CellRange

from api.specifications.table_structure import (
    SourceLocator,
    TableLimits,
    TableStructure,
    TableStructureError,
    table_from_hwpx_xml,
    table_from_xlsx,
)

MAX_BYTES = 50_000_000
MAX_ENTRIES = 500
MAX_TEXT = 200_000
MAX_CELLS = 100_000
MAX_TABLES = 100
MAX_TABLE_METADATA = 2_000_000
HEADERS = {
    "item": {"품명", "품목", "품목명", "물품명", "item"},
    "specification": {"규격", "사양", "specification"},
    "quantity": {"수량", "quantity", "qty"},
    "unit": {"단위", "unit"},
    "unit_price": {"단가", "unitprice", "unit_price"},
    "amount": {"금액", "합계금액", "amount"},
}


class ExtractionError(ValueError):
    pass


class TableCollector:
    """Additive table metadata; failures never replace legacy text/completeness."""

    def __init__(self) -> None:
        self.tables: list[dict] = []
        self.issues: list[dict] = []
        self.unknown = False
        self.detected = 0
        self.size = 0
        self.cells = 0
        self.omitted_issues = 0

    def reject(self, code: str, source: SourceLocator) -> None:
        issue = {"code": code, "source": source.model_dump(mode="json")}
        if issue not in self.issues:
            if len(self.issues) < MAX_TABLES:
                self.issues.append(issue)
            else:
                self.omitted_issues += 1

    def add(self, table: TableStructure) -> None:
        data = table.model_dump(mode="json")
        size = 0
        for part in json.JSONEncoder(ensure_ascii=False).iterencode(data):
            size += len(part)
            if self.size + size > MAX_TABLE_METADATA:
                raise TableStructureError("table_metadata_limit")
        if len(self.tables) >= MAX_TABLES or self.cells + len(table.cells) > MAX_CELLS:
            raise TableStructureError("table_collection_limit")
        self.tables.append(data)
        self.cells += len(table.cells)
        self.size += size

    def metadata(self) -> dict:
        review = bool(self.issues) or any(t["needs_review"] for t in self.tables)
        status = "needs_review" if review else "preserved" if self.tables else "not_detected"
        if self.issues:
            status = "partial" if self.tables else "needs_review"
        elif self.unknown:
            status = "partial" if self.tables else "unknown"
        return {"tables": self.tables, "table_structure": {
            "version": 1, "status": status, "needs_review": review,
            "detected_count": None if self.unknown else self.detected,
            "preserved_count": len(self.tables), "issues": self.issues,
            "omitted_issue_count": self.omitted_issues,
            "structure_trustworthy": bool(self.tables) and not review and not self.unknown,
            "source_locations": [table["source"] for table in self.tables],
        }}


def text_only_table_metadata(*, locator: str, table_count: int | None = None) -> dict:
    """Text-only HWP cannot establish table presence, let alone geometry."""
    collector = TableCollector()
    collector.unknown = True
    if type(table_count) is int and table_count > 0:
        collector.unknown = False
        collector.detected = table_count
        collector.reject("hwp_table_structure_unavailable", SourceLocator(locator=locator))
    return collector.metadata()


def _xlsx_structure(sheet: object, merges: list[str], values: list[tuple], locator: str) -> TableStructure:
    limits = TableLimits()
    rows, cols = sheet.max_row or 1, sheet.max_column or 1
    if len(merges) > limits.max_cells:
        raise TableStructureError("table_span_limit")
    ranges = [CellRange(ref) for ref in merges]
    rows = max([rows, *(r.max_row for r in ranges)])
    cols = max([cols, *(r.max_col for r in ranges)])
    if rows > limits.max_rows or cols > limits.max_cols or rows * cols > limits.max_cells:
        raise TableStructureError("table_dimension_limit")
    if sum(r.size["rows"] * r.size["columns"] for r in ranges) > limits.max_cells:
        raise TableStructureError("table_span_limit")
    book = Workbook()
    try:
        target = book.active
        target.title = sheet.title
        target.cell(rows, cols)
        for row, col, value, data_type in values:
            target.cell(row, col, value).data_type = data_type
        # Register explicit ranges without allocating or erasing covered cells.
        interiors = set()
        for region in ranges:
            target.merged_cells.add(region)
            interiors.update((row, col) for row in range(region.min_row, region.max_row + 1)
                             for col in range(region.min_col, region.max_col + 1)
                             if (row, col) != (region.min_row, region.min_col))
        if any((row, col) in interiors and value is not None for row, col, value, _ in values):
            raise TableStructureError("merged_non_anchor_content")
        bands, label_cols = [], set()
        skip_until = 0
        for row in target.iter_rows():
            if row[0].row <= skip_until:
                continue
            detected = {}
            for cell in row:
                label = str(cell.value or "").lower().replace(" ", "")
                for kind, aliases in HEADERS.items():
                    if label in aliases:
                        detected[cell.column] = kind
            item_cols = [col for col, kind in detected.items() if kind == "item"]
            if len(item_cols) != 1:
                continue
            item_col = item_cols[0]
            start = row[0].row
            end = max([start, *(r.max_row for r in ranges if r.min_row == start and r.min_col == item_col)])
            if len(detected) >= 2 or end > start:
                if end - start >= 32:
                    raise TableStructureError("table_header_limit")
                bands.append(tuple(range(start - 1, end)))
                label_cols.add(item_col - 1)
                skip_until = end
        return table_from_xlsx(target, locator=locator, header_bands=tuple(bands),
                               row_label_col=next(iter(label_cols)) if len(label_cols) == 1 else None)
    finally:
        book.close()


def checked_archive(content: bytes, budget: list[int] | None = None) -> zipfile.ZipFile:
    if len(content) > MAX_BYTES:
        raise ExtractionError("archive_size_limit")
    archive = zipfile.ZipFile(io.BytesIO(content))
    try:
        entries = archive.infolist()
        if len(entries) > MAX_ENTRIES or sum(e.file_size for e in entries) > MAX_BYTES:
            raise ExtractionError("archive_size_limit")
        if budget is not None:
            budget[0] -= sum(e.file_size for e in entries)
            if budget[0] < 0:
                raise ExtractionError("archive_size_limit")
        names = set()
        for entry in entries:
            name = entry.filename
            path = PurePosixPath(name)
            if (path.is_absolute() or ".." in path.parts or "\\" in name or ":" in name
                or any(ord(c) < 32 for c in name) or entry.orig_filename != name
                or stat.S_ISLNK(entry.external_attr >> 16) or entry.flag_bits & 1):
                raise ExtractionError("unsafe_archive_entry")
            identity = name.casefold().rstrip("/")
            if identity in names:
                raise ExtractionError("ambiguous_archive_entry")
            names.add(identity)
            if entry.file_size > max(entry.compress_size, 1) * 500:
                raise ExtractionError("archive_expansion_limit")
        return archive
    except Exception:
        archive.close()
        raise


def _value(value: object) -> str | int | float | bool | None:
    if isinstance(value, date | datetime):
        return value.isoformat()
    if value is None or type(value) in (int, float, bool, str):
        return value
    return str(value)


def extract_xlsx(content: bytes, *, expansion_budget: list[int] | None = None,
                 locator: str | None = None) -> dict:
    locator = locator or f"sha256:{hashlib.sha256(content).hexdigest()}/xlsx"
    merge_parts: dict[str, list[str]] = {}
    with checked_archive(content, expansion_budget) as archive:
        if "xl/workbook.xml" not in archive.namelist():
            raise ExtractionError("invalid_spreadsheet")
        for entry in archive.infolist():
            if entry.filename.lower().endswith(".xml"):
                try:
                    root = fromstring(archive.read(entry), forbid_dtd=True)
                    if root.tag.rsplit("}", 1)[-1] == "worksheet":
                        merge_parts[entry.filename] = [
                            element.attrib.get("ref", "") for element in root.iter()
                            if element.tag.rsplit("}", 1)[-1] == "mergeCell"
                        ]
                except Exception:
                    raise ExtractionError("unsafe_spreadsheet_xml") from None
    options = {"read_only": True, "keep_links": False}
    formulas = load_workbook(io.BytesIO(content), data_only=False, **options)
    cached = None
    try:
        cached = load_workbook(io.BytesIO(content), data_only=True, **options)
        lines, cells, items, issues = [], [], [], []
        tables = TableCollector()
        visited, text_size = 0, 0
        for sheet in formulas:
            structure_values = []
            if (sheet.max_row or 0) > 10_000 or (sheet.max_column or 0) > 200:
                raise ExtractionError("spreadsheet_cell_limit")
            columns = {}
            for row, cached_row in zip(sheet.iter_rows(), cached[sheet.title].iter_rows(), strict=True):
                visited += len(row)
                if visited > MAX_CELLS:
                    raise ExtractionError("spreadsheet_cell_limit")
                records, detected = {}, {}
                for cell, cache in zip(row, cached_row, strict=True):
                    if cell.value is None:
                        continue
                    structure_values.append((cell.row, cell.column, cell.value, cell.data_type))
                    value = _value(cache.value if cell.data_type == "f" else cell.value)
                    record = {"sheet": sheet.title, "cell": cell.coordinate, "value": value,
                              "hidden_sheet": sheet.sheet_state != "visible",
                              "formula": str(cell.value) if cell.data_type == "f" else None}
                    if cell.data_type in {"f", "e"}:
                        issues.append({"code": "formula_review" if cell.data_type == "f" else "cell_error",
                                       "sheet": sheet.title, "cell": cell.coordinate})
                    cells.append(record)
                    records[cell.column] = record
                    label = str(value or "").lower().replace(" ", "")
                    for kind, aliases in HEADERS.items():
                        if label in aliases:
                            detected[cell.column] = kind
                    if value is not None:
                        line = f"[시트 {sheet.title} · 셀 {cell.coordinate}] {value}"
                        text_size += len(line) + 1
                        if text_size > MAX_TEXT:
                            raise ExtractionError("extracted_text_limit")
                        lines.append(line)
                if "item" in detected.values() and len(detected) >= 2:
                    columns = detected
                elif columns:
                    mapped = {kind: records[column] for column, kind in columns.items() if column in records}
                    if mapped.get("item", {}).get("value") is not None:
                        items.append({"sheet": sheet.title, "row": mapped["item"]["cell"],
                                      "fields": mapped, "review_status": "pending"})
            if structure_values:
                tables.detected += 1
                source = SourceLocator(locator=f"{locator}#sheet={sheet.title}")
                try:
                    if tables.detected > MAX_TABLES:
                        raise TableStructureError("table_collection_limit")
                    # ReadOnlyWorksheet retains the resolved OOXML part path; do
                    # not guess sheetN.xml from display order or load merge cells.
                    part = sheet._worksheet_path.lstrip("/")
                    if part not in merge_parts:
                        raise TableStructureError("xlsx_geometry_part_missing")
                    tables.add(_xlsx_structure(sheet, merge_parts[part], structure_values, locator))
                except TableStructureError as exc:
                    tables.reject(str(exc), source)
                except (ValueError, TypeError, AttributeError):
                    tables.reject("xlsx_table_structure_unavailable", source)
        return {"text": "\n".join(lines), "format": "xlsx", "extraction": {
            "version": 1, "complete": bool(lines) and not issues, "cells": cells,
            "items": items, "issues": issues, "formula_evaluation": "not_performed",
            **tables.metadata(),
        }}
    finally:
        formulas.close()
        if cached is not None:
            cached.close()


def extract_hwpx(content: bytes, *, locator: str | None = None) -> dict:
    """Keep legacy itertext order exactly; attach XML-addressed tables separately."""
    locator = locator or f"sha256:{hashlib.sha256(content).hexdigest()}/hwpx"
    texts, length = [], 0
    tables = TableCollector()
    with checked_archive(content) as archive:
        entries = [e for e in archive.infolist()
                   if e.filename.startswith("Contents/section") and e.filename.endswith(".xml")]
        for entry in entries:
            xml = archive.read(entry)
            root = fromstring(xml, forbid_dtd=True, forbid_entities=True, forbid_external=True)
            text = " ".join(root.itertext())
            length += len(text) + bool(texts)
            if length > MAX_TEXT:
                raise ExtractionError("extracted_text_limit")
            texts.append(text)
            count = sum(element.tag.rsplit("}", 1)[-1] == "tbl" for element in root.iter())
            base = f"{locator}/{entry.filename}"
            for index in range(count):
                tables.detected += 1
                source = SourceLocator(locator=f"{base}#table={index}")
                try:
                    if tables.detected > MAX_TABLES:
                        raise TableStructureError("table_collection_limit")
                    tables.add(table_from_hwpx_xml(xml, locator=base, table_index=index))
                except TableStructureError as exc:
                    tables.reject(str(exc), source)
                except (ValueError, TypeError):
                    tables.reject("hwpx_table_structure_unavailable", source)
    return {"text": "\n".join(texts), "format": "hwpx", "extraction": {
        "version": 1, "complete": bool(texts) and any(t.strip() for t in texts),
        **tables.metadata(),
    }}


def extract_zip(content: bytes, *, ocr_budget: list[int] | None = None) -> dict:
    from urllib.parse import quote

    from api.workspace.pdf_text import extract_pdf

    members, texts, total = [], [], 0
    expansion_budget = [MAX_BYTES]
    ocr_budget = [10] if ocr_budget is None else ocr_budget
    tables = TableCollector()
    base = f"sha256:{hashlib.sha256(content).hexdigest()}/zip"
    with checked_archive(content, expansion_budget) as archive:
        for entry in archive.infolist():
            if entry.is_dir():
                continue
            data = archive.read(entry)
            member = {"name": entry.filename, "sha256": hashlib.sha256(data).hexdigest(),
                      "size": len(data), "status": "unsupported", "complete": False}
            suffix = PurePosixPath(entry.filename).suffix.lower()
            try:
                if suffix == ".xlsx" and data.startswith(b"PK"):
                    result = extract_xlsx(data, expansion_budget=expansion_budget,
                                          locator=f"{base}/{quote(entry.filename, safe='')}")
                elif suffix == ".pdf" and data.startswith(b"%PDF"):
                    result = extract_pdf(data, ocr_budget=ocr_budget,
                                         locator=f"{base}/{quote(entry.filename, safe='')}")
                else:
                    tables.unknown = True
                    members.append(member)
                    continue
                member.update(status="extracted" if result["text"].strip() else "unread",
                              complete=result["extraction"]["complete"], extraction=result["extraction"])
                table_meta = result["extraction"]["table_structure"]
                tables.detected += table_meta["detected_count"] or 0
                tables.unknown |= table_meta["detected_count"] is None
                for issue in table_meta["issues"]:
                    tables.reject(issue["code"], SourceLocator.model_validate(issue["source"]))
                for table in result["extraction"]["tables"]:
                    try:
                        tables.add(TableStructure.model_validate(table))
                    except ValueError:
                        tables.reject("table_collection_limit", SourceLocator.model_validate(table["source"]))
                if result["text"].strip():
                    block = f"[ZIP 파일 {entry.filename}]\n{result['text']}"
                    total += len(block) + 2
                    if total > MAX_TEXT:
                        raise ExtractionError("extracted_text_limit")
                    texts.append(block)
            except ExtractionError:
                raise
            except Exception:
                tables.unknown = True
                member.update(status="failed", error_code="member_extraction_failed")
            members.append(member)
    return {"text": "\n\n".join(texts), "format": "zip", "extraction": {
        "version": 1, "complete": bool(members) and all(m["complete"] for m in members),
        "members": members,
        **tables.metadata(),
    }}
