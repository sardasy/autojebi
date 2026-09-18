"""Excel job: requirement/evidence workbook for review."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from api.workspace.service import blockers
from api.workspace.storage import artifact


def run_excel(job: dict, folder: Path) -> dict:
    book = Workbook()
    sheet = book.active
    sheet.title = "규격 및 근거"
    sheet.append(
        ["요구사항", "원문 인용", "원문 위치", "제안값", "근거 인용", "근거 위치", "검토 상태"]
    )
    for r in job["snapshot"]["requirements"]:
        values = [
            r.get(k, "")
            for k in (
                "text",
                "quote",
                "locator",
                "proposed_value",
                "evidence_quote",
                "evidence_locator",
                "review_status",
            )
        ]
        # External document text must not become spreadsheet formulas.
        sheet.append(
            [
                "'" + v if isinstance(v, str) and v.startswith(("=", "+", "-", "@")) else v
                for v in values
            ]
        )
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for col in "ABCDEFG":
        sheet.column_dimensions[col].width = 35 if col in "ABDE" else 22
    for row in sheet:
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        sheet.row_dimensions[row[0].row].height = 60
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="444444")
    summary = book.create_sheet("검토 기록")
    summary.append(["공고번호", job["notice_no"]])
    summary.append(["입력 버전", job["input_version"]])
    warnings = blockers(job["snapshot"])
    for warning in warnings:
        summary.append(["확인 필요", warning])
    summary.column_dimensions["A"].width = 20
    summary.column_dimensions["B"].width = 95
    output = folder / "requirements.xlsx"
    book.save(output)
    return artifact(output, warnings)
