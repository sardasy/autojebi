from __future__ import annotations

import json
from xml.sax.saxutils import escape

import pytest
from openpyxl import Workbook
from pydantic import ValidationError

from api.specifications.table_structure import (
    ColumnUnit,
    SourceLocator,
    SourceText,
    TableCell,
    TableLimits,
    TableStructure,
    TableStructureError,
    build_table_chunks,
    context_for_span,
    preserve_table_structure,
    table_from_hwpx_xml,
    table_from_pdf_cells,
    table_from_xlsx,
)


def source(text: str, locator: str = "fixture:title") -> SourceText:
    return SourceText(raw_text=text, source=SourceLocator(locator=locator))


def cell(row: int, col: int, text: str, **kwargs: object) -> TableCell:
    return TableCell(
        row=row,
        col=col,
        raw_text=text,
        source=SourceLocator(
            locator=f"fixture:r{row}c{col}",
            page=1,
            bbox=(col * 20, row * 10, col * 20 + 20, row * 10 + 10),
        ),
        **kwargs,
    )


@pytest.fixture
def worksheet():
    book = Workbook()
    sheet = book.active
    sheet.title = "Specifications"
    sheet.append(["Item", "Rating", None])
    sheet.append([None, "Voltage", "Current"])
    sheet.append(["Module", " 400\nnominal ", 0])
    sheet.append([None, None, False])
    sheet.merge_cells("A1:A2")
    sheet.merge_cells("B1:C1")
    yield sheet
    book.close()


def hwpx(cells: list[tuple[int, int, int, int, str]], rows: int, cols: int) -> bytes:
    # Physical XML order deliberately differs from logical address order.
    content = "".join(
        f"<hp:tc><hp:subList><hp:p><hp:run><hp:t>{escape(text)}</hp:t>"
        f'</hp:run></hp:p></hp:subList><hp:cellAddr rowAddr="{row}" colAddr="{col}"/>'
        f'<hp:cellSpan rowSpan="{height}" colSpan="{width}"/></hp:tc>'
        for row, col, height, width, text in cells
    )
    return (
        '<hp:sec xmlns:hp="http://www.hancom.co.kr/hwpml/2011/paragraph">'
        f'<hp:tbl rowCnt="{rows}" colCnt="{cols}"><hp:tr>{content}</hp:tr>'
        "</hp:tbl></hp:sec>"
    ).encode()


def by_position(table: TableStructure) -> dict[tuple[int, int], TableCell]:
    return {(c.row, c.col): c for c in table.cells}


def test_xlsx_merged_two_level_headers_and_context(worksheet):
    title = source("Power module specifications")
    units = (ColumnUnit(col=1, value=source("V", "fixture:unit-note")),)
    table = table_from_xlsx(
        worksheet, locator="document:42", title=title, header_bands=((0, 1),), units=units
    )
    assert not table.needs_review
    cells = by_position(table)
    assert len(cells) == 10
    assert cells[0, 0].row_span == 2
    assert cells[0, 1].col_span == 2
    assert (0, 2) not in cells  # Covered position is not a copied value.
    assert [h.raw_text for h in cells[2, 1].header_paths[0].headers] == ["Rating", "Voltage"]
    assert [h.raw_text for h in cells[2, 0].header_paths[0].headers] == ["Item"]
    assert cells[2, 1].header_paths[0].headers[0].source.locator.endswith("!B1:C1")
    assert cells[2, 1].raw_text == " 400\nnominal "
    assert cells[2, 2].raw_text == "0"
    chunks = build_table_chunks(table)
    chunk = next(c for c in chunks if c.row == 2)
    payload = json.loads(chunk.text)
    assert payload["title"]["raw_text"] == title.raw_text
    assert payload["row_label"]["raw_text"] == "Module"
    assert payload["units"][0]["value"]["raw_text"] == "V"
    assert title.source in chunk.source_locators
    assert units[0].value.source in chunk.source_locators
    assert cells[0, 1].source in chunk.source_locators
    assert sum(len(json.loads(c.text)["cells"]) for c in chunks) == len(table.cells)
    assert table == preserve_table_structure(table)
    assert chunks == build_table_chunks(table)
    assert TableStructure.model_validate_json(table.model_dump_json()) == table


def test_xlsx_empty_cells_never_forward_filled(worksheet):
    table = table_from_xlsx(worksheet, locator="book", header_bands=((0, 1),))
    cells = by_position(table)
    assert cells[3, 0].raw_text == cells[3, 1].raw_text == ""
    assert cells[3, 2].raw_text == "False"
    last = json.loads(build_table_chunks(table)[-1].text)
    assert last["row_label"]["raw_text"] == ""
    assert last["row_label"]["source"]["locator"].endswith("!A4")


def test_hwpx_explicit_spans_match_xlsx_and_keep_text():
    xml = hwpx(
        [
            (2, 2, 1, 1, "20"),
            (0, 1, 1, 2, "Rating"),
            (0, 0, 2, 1, "Item"),
            (1, 1, 1, 1, "Voltage"),
            (1, 2, 1, 1, "Current"),
            (2, 0, 1, 1, "Module"),
            (2, 1, 1, 1, " 400 & 500\n "),
        ],
        3,
        3,
    )
    table = table_from_hwpx_xml(
        xml, locator="doc:42/Contents/section0.xml", header_bands=((0, 1),), page=7
    )
    assert not table.needs_review
    cells = by_position(table)
    assert cells[0, 0].row_span == 2
    assert cells[0, 1].col_span == 2
    assert cells[2, 1].raw_text == " 400 & 500\n "
    assert [h.raw_text for h in cells[2, 1].header_paths[0].headers] == ["Rating", "Voltage"]
    assert cells[2, 1].source.page == 7
    assert cells[2, 1].source.bbox is None
    assert "cell=6;r=2;c=1" in cells[2, 1].source.locator


@pytest.mark.parametrize("format", ["hwpx", "xlsx", "pdf"])
def test_conflicting_spans_retained_and_never_used_for_headers(format):
    if format == "hwpx":
        table = table_from_hwpx_xml(
            hwpx([(0, 0, 1, 2, "Left"), (0, 1, 1, 1, "Right"), (1, 0, 1, 1, "Module")], 2, 2),
            locator="doc",
            header_bands=((0,),),
        )
    elif format == "xlsx":
        sheet = Workbook().active
        sheet.append(["Left", None, None])
        sheet.append(["Module", "1", "2"])
        sheet.merge_cells("A1:B1")
        sheet.merge_cells("B1:C1")
        table = table_from_xlsx(sheet, locator="doc", header_bands=((0,),))
    else:
        table = table_from_pdf_cells(
            [cell(0, 0, "Left", col_span=2), cell(0, 1, "Right"), cell(1, 0, "Module")],
            locator="doc",
            n_rows=2,
            n_cols=2,
            header_bands=((0,),),
            headers_confirmed=True,
        )
    assert table.needs_review
    assert "conflicting_spans" in {i.code for i in table.issues}
    assert sum(c.row == 0 for c in table.cells) == 2
    assert all(not c.header_paths for c in table.cells)
    chunks = build_table_chunks(table)
    assert all(c.needs_review for c in chunks)
    assert sum(len(json.loads(c.text)["cells"]) for c in chunks) == len(table.cells)


@pytest.mark.parametrize("confirmed,ambiguous", [(False, False), (True, True)])
def test_ambiguous_pdf_ocr_has_no_inferred_headers(confirmed, ambiguous):
    original = [cell(0, 0, "Item"), cell(0, 1, "Volt?"), cell(1, 0, ""), cell(1, 1, "4OO")]
    table = table_from_pdf_cells(
        original,
        locator="scan",
        n_rows=2,
        n_cols=2,
        header_bands=((0,),),
        headers_confirmed=confirmed,
        ocr_ambiguous=ambiguous,
    )
    assert table.needs_review
    assert all(not c.header_paths for c in table.cells)
    assert [c.raw_text for c in table.cells] == [c.raw_text for c in original]
    last = json.loads(build_table_chunks(table)[-1].text)
    assert last["row_label"]["raw_text"] == ""
    assert last["cells"][1]["source"]["bbox"] == [20.0, 10.0, 40.0, 20.0]
    assert last["cells"][1]["source"]["page"] == 1


def test_blank_pdf_header_is_not_filled_from_neighbor():
    table = table_from_pdf_cells(
        [cell(0, 0, "Item"), cell(0, 1, ""), cell(1, 0, "Module"), cell(1, 1, "10")],
        locator="pdf",
        n_rows=2,
        n_cols=2,
        header_bands=((0,),),
        headers_confirmed=True,
    )
    assert table.needs_review
    assert by_position(table)[1, 1].header_paths[0].headers == ()
    assert by_position(table)[0, 1].raw_text == ""


def test_repeated_headers_use_their_own_page_locators():
    cells = [
        cell(r, c, text)
        for r, texts in enumerate(
            [("Item", "Voltage"), ("Module A", "100"), ("Item", "Voltage"), ("Module B", "200")]
        )
        for c, text in enumerate(texts)
    ]
    cells = [
        c.model_copy(update={"source": c.source.model_copy(update={"page": 1 + c.row // 2})})
        for c in cells
    ]
    table = table_from_pdf_cells(
        cells, locator="pdf", n_rows=4, n_cols=2, header_bands=((0,), (2,)), headers_confirmed=True
    )
    assert not table.needs_review
    chunks = build_table_chunks(table)
    assert [c.kind for c in chunks] == ["header", "data", "header", "data"]
    path = by_position(table)[3, 1].header_paths[0]
    assert path.headers[0].source == cells[5].source
    assert path.headers[0].source.page == 2
    assert cells[5].source in chunks[-1].source_locators
    assert cells[1].source not in chunks[-1].source_locators


def test_header_data_span_requires_review():
    table = table_from_pdf_cells(
        [cell(0, 0, "Item", row_span=2), cell(0, 1, "Voltage"), cell(1, 1, "200")],
        locator="pdf",
        n_rows=2,
        n_cols=2,
        header_bands=((0,),),
        headers_confirmed=True,
    )
    assert table.needs_review
    assert "header_data_span" in {i.code for i in table.issues}
    assert not by_position(table)[1, 1].header_paths


def test_merged_value_preserves_both_header_paths_but_requires_review():
    table = table_from_pdf_cells(
        [cell(0, 0, "Min"), cell(0, 1, "Max"), cell(1, 0, "10", col_span=2)],
        locator="pdf",
        n_rows=2,
        n_cols=2,
        header_bands=((0,),),
        headers_confirmed=True,
    )
    assert "merged_value_multiple_headers" in {i.code for i in table.issues}
    assert len(by_position(table)[1, 0].header_paths) == 2


def test_formula_not_evaluated_or_silently_replaced(worksheet):
    worksheet["B3"] = "=1+2"
    table = table_from_xlsx(worksheet, locator="book", header_bands=((0, 1),))
    assert table.needs_review
    assert by_position(table)[2, 1].raw_text == "=1+2"
    assert "formula_or_error_review" in {i.code for i in table.issues}


def test_missing_cells_are_not_synthesized():
    table = table_from_pdf_cells(
        [cell(0, 0, "Item"), cell(1, 1, "10")], locator="pdf", n_rows=2, n_cols=2
    )
    assert len(table.cells) == 2
    assert "missing_cells" in {i.code for i in table.issues}
    assert json.loads(build_table_chunks(table)[1].text)["row_label"] is None


@pytest.mark.parametrize(
    "limits,code",
    [
        (TableLimits(max_rows=3), "table_dimension_limit"),
        (TableLimits(max_cols=2), "table_dimension_limit"),
        (TableLimits(max_cells=11), "table_cell_limit"),
        (TableLimits(max_text_chars=10), "table_text_limit"),
    ],
)
def test_xlsx_bounds_are_errors_not_truncation(worksheet, limits, code):
    with pytest.raises(TableStructureError, match=code):
        table_from_xlsx(worksheet, locator="book", limits=limits)


@pytest.mark.parametrize(
    "limits,code",
    [
        (TableLimits(max_xml_bytes=20), "xml_byte_limit"),
        (TableLimits(max_xml_nodes=2), "xml_node_limit"),
        (TableLimits(max_text_chars=3), "table_text_limit"),
    ],
)
def test_hwpx_bounds_are_errors_not_truncation(limits, code):
    with pytest.raises(TableStructureError, match=code):
        table_from_hwpx_xml(hwpx([(0, 0, 1, 1, "Hello")], 1, 1), locator="xml", limits=limits)


@pytest.mark.parametrize(
    "xml",
    [
        b'<!DOCTYPE x [<!ENTITY x "unsafe">]><x>&x;</x>',
        b'<!DOCTYPE x SYSTEM "file:///private"><x/>',
        b"<x>",
    ],
)
def test_hwpx_unsafe_xml_rejected(xml):
    with pytest.raises(TableStructureError, match="unsafe_or_invalid_hwpx_xml"):
        table_from_hwpx_xml(xml, locator="xml")


@pytest.mark.parametrize(
    "height,width,code",
    [
        (0, 1, "invalid_hwpx_geometry"),
        (-1, 1, "invalid_hwpx_geometry"),
        (1, 2, "span_out_of_bounds"),
        (1_000_000, 1, "span_out_of_bounds"),
    ],
)
def test_hwpx_invalid_span(height, width, code):
    with pytest.raises(TableStructureError, match=code):
        table_from_hwpx_xml(hwpx([(0, 0, height, width, "X")], 1, 1), locator="xml")


def test_hwpx_requires_explicit_geometry_and_preserves_empty_cell():
    xml = hwpx([(0, 0, 1, 1, "")], 1, 1)
    table = table_from_hwpx_xml(xml, locator="xml", header_bands=((0,),))
    assert len(table.cells) == 1 and table.cells[0].raw_text == ""
    with pytest.raises(TableStructureError, match="missing_or_conflicting_hwpx_geometry"):
        table_from_hwpx_xml(
            xml.replace(b'<hp:cellSpan rowSpan="1" colSpan="1"/>', b""), locator="xml"
        )


@pytest.mark.parametrize("bands", [((1, 0),), ((0,), (0,)), ((0, 2),), ((4,),)])
def test_invalid_header_bands_rejected(worksheet, bands):
    with pytest.raises(TableStructureError, match="invalid_header_bands"):
        table_from_xlsx(worksheet, locator="book", header_bands=bands)


@pytest.mark.parametrize(
    "limits,code",
    [
        (TableLimits(max_chunk_chars=100), "table_chunk_size_limit"),
        (TableLimits(max_chunks=1), "table_chunk_count_limit"),
        (TableLimits(max_output_chars=100), "table_chunk_output_limit"),
    ],
)
def test_chunk_limits_never_return_partial_results(worksheet, limits, code):
    table = table_from_xlsx(worksheet, locator="book", header_bands=((0, 1),))
    before = table.model_dump_json()
    with pytest.raises(TableStructureError, match=code):
        build_table_chunks(table, limits=limits)
    assert table.model_dump_json() == before


def test_exact_chunk_limits_and_escaping_are_lossless(worksheet):
    worksheet["B3"] = 'value | "header"\n\tend'
    table = table_from_xlsx(worksheet, locator="book", header_bands=((0, 1),))
    chunks = build_table_chunks(table)
    maximum = max(len(c.text) for c in chunks)
    total = sum(len(c.text) for c in chunks)
    limits = TableLimits(max_chunk_chars=maximum, max_output_chars=total, max_chunks=len(chunks))
    assert chunks == build_table_chunks(table, limits=limits)
    assert json.loads(chunks[2].text)["cells"][1]["raw_text"] == worksheet["B3"].value
    with pytest.raises(TableStructureError, match="table_chunk_size_limit"):
        build_table_chunks(table, limits=limits.model_copy(update={"max_chunk_chars": maximum - 1}))


@pytest.mark.parametrize("kwargs", [{"row_span": 0}, {"col_span": -1}, {"row_span": True}])
def test_cell_models_reject_invalid_spans(kwargs):
    with pytest.raises(ValidationError):
        cell(0, 0, "x", **kwargs)


@pytest.mark.parametrize("bbox", [(2, 0, 1, 1), (0, 2, 1, 1), (0, 0, float("nan"), 1)])
def test_locator_rejects_invalid_bbox(bbox):
    with pytest.raises(ValidationError):
        SourceLocator(locator="pdf", page=1, bbox=bbox)


def test_pdf_span_work_budget_checked_before_expansion():
    with pytest.raises(TableStructureError, match="table_span_limit"):
        table_from_pdf_cells(
            [cell(0, 0, "a", col_span=2), cell(0, 0, "b", col_span=2)],
            locator="pdf",
            n_rows=1,
            n_cols=2,
            limits=TableLimits(max_cells=3),
        )


def test_sparse_xlsx_dimension_checked_before_materialization():
    sheet = Workbook().active
    sheet["A1048576"] = "sparse"
    with pytest.raises(TableStructureError, match="table_dimension_limit"):
        table_from_xlsx(sheet, locator="book")


def test_hwpx_text_controls_and_multiple_paragraphs_keep_order():
    xml = hwpx([(0, 0, 1, 1, "REPLACE")], 1, 1)
    xml = xml.replace(b"REPLACE", b"before<hp:tab/>middle<hp:lineBreak/>after")
    xml = xml.replace(b"</hp:p>", b"</hp:p><hp:p><hp:run><hp:t>next</hp:t></hp:run></hp:p>")
    table = table_from_hwpx_xml(xml, locator="xml")
    assert table.cells[0].raw_text == "before\tmiddle\nafter\nnext"


def test_hwpx_nested_table_is_explicitly_rejected():
    xml = hwpx([(0, 0, 1, 1, "REPLACE")], 1, 1)
    xml = xml.replace(b"REPLACE", b'<hp:tbl rowCnt="1" colCnt="1"/>')
    with pytest.raises(TableStructureError, match="nested_hwpx_table_requires_review"):
        table_from_hwpx_xml(xml, locator="xml")


def test_span_across_repeated_header_bands_is_not_assigned():
    table = table_from_pdf_cells(
        [cell(0, 0, "Header", row_span=2), cell(2, 0, "10")],
        locator="pdf",
        n_rows=3,
        n_cols=1,
        header_bands=((0,), (1,)),
        headers_confirmed=True,
    )
    assert table.needs_review
    assert "header_band_span" in {i.code for i in table.issues}
    assert not by_position(table)[2, 0].header_paths


def test_pdf_missing_geometry_is_review_only():
    original = cell(0, 0, "Header").model_copy(update={"source": SourceLocator(locator="pdf:cell")})
    table = table_from_pdf_cells(
        [original], locator="pdf", n_rows=1, n_cols=1, header_bands=((0,),), headers_confirmed=True
    )
    assert "missing_pdf_geometry" in {i.code for i in table.issues}
    assert table.needs_review


def test_out_of_bounds_and_conflicting_units_are_rejected(worksheet):
    for units, code in [
        ((ColumnUnit(col=3, value=source("V")),), "unit_out_of_bounds"),
        (
            (ColumnUnit(col=1, value=source("V")), ColumnUnit(col=1, value=source("A"))),
            "conflicting_units",
        ),
    ]:
        with pytest.raises(TableStructureError, match=code):
            table_from_xlsx(worksheet, locator="book", units=units)


def test_xlsx_limit_equality_accepts_all_cells(worksheet):
    table = table_from_xlsx(worksheet, locator="book", header_bands=((0, 1),))
    limits = TableLimits(
        max_rows=4,
        max_cols=3,
        max_cells=12,
        max_text_chars=sum(len(c.raw_text) for c in table.cells),
    )
    assert (
        table_from_xlsx(worksheet, locator="book", header_bands=((0, 1),), limits=limits) == table
    )


def test_direct_table_bounds_revalidated_by_chunk_builder():
    table = TableStructure(
        format="pdf",
        source=SourceLocator(locator="pdf"),
        n_rows=1,
        n_cols=1,
        cells=(cell(0, 0, "Value", col_span=2),),
    )
    with pytest.raises(TableStructureError, match="span_out_of_bounds"):
        build_table_chunks(table)


def test_row_label_context_uses_only_explicit_vertical_merge():
    sheet = Workbook().active
    sheet.append(["Item", "Value"])
    sheet.append(["Module", "100"])
    sheet.append([None, "200"])
    sheet.append([None, "300"])
    sheet.merge_cells("A2:A3")
    table = table_from_xlsx(sheet, locator="book", header_bands=((0,),))
    chunks = build_table_chunks(table)
    explicit = json.loads(chunks[2].text)
    assert explicit["row_label"]["raw_text"] == "Module"
    assert explicit["row_label"]["source"]["locator"].endswith("!A2:A3")
    assert by_position(table)[1, 0].source in chunks[2].source_locators
    assert len(explicit["cells"]) == 1
    assert json.loads(chunks[3].text)["row_label"]["raw_text"] == ""


def test_empty_header_only_pdf_still_requires_review():
    table = table_from_pdf_cells(
        [cell(0, 0, "")],
        locator="pdf",
        n_rows=1,
        n_cols=1,
        header_bands=((0,),),
        headers_confirmed=True,
    )
    assert table.needs_review
    assert "ambiguous_headers" in {issue.code for issue in table.issues}


def span_extraction(*tables: TableStructure) -> dict:
    review = any(table.needs_review for table in tables)
    return {
        "tables": [table.model_dump(mode="json") for table in tables],
        "table_structure": {
            "version": 1,
            "status": "needs_review" if review else "preserved",
            "needs_review": review,
            "detected_count": len(tables),
            "preserved_count": len(tables),
            "issues": [],
            "omitted_issue_count": 0,
        },
    }


@pytest.fixture
def context_table():
    sheet = Workbook().active
    sheet.append(["Item", "Electrical", None])
    sheet.append([None, "Voltage", "Current"])
    sheet.append(["Module X", "400 V", "20 A"])
    sheet.append([None, "240 V", "30 A"])
    sheet.merge_cells("A1:A2")
    sheet.merge_cells("B1:C1")
    sheet.merge_cells("A3:A4")
    return table_from_xlsx(
        sheet,
        locator="fixture:document",
        header_bands=((0, 1),),
        title=source("Power module", "fixture:title"),
        units=(
            ColumnUnit(col=1, value=source("V", "fixture:unit-V")),
            ColumnUnit(col=2, value=source("A", "fixture:unit-A")),
        ),
    )


def test_span_context_attaches_only_unique_cell_and_its_context(context_table):
    extraction = span_extraction(context_table)
    before = json.dumps(extraction)
    text = "Unrelated\n240 V\nEnd"
    result = context_for_span(extraction, text, 10, 15)
    assert result["status"] == "matched" and not result["needs_review"]
    contexts = json.loads(result["text"])
    assert len(contexts) == 1
    assert contexts[0]["cell"]["raw_text"] == "240 V"
    assert contexts[0]["title"]["raw_text"] == "Power module"
    assert contexts[0]["row_label"]["raw_text"] == "Module X"
    assert contexts[0]["row_label"]["source"]["locator"].endswith("!A3:A4")
    assert [h["raw_text"] for h in contexts[0]["header_paths"][0]["headers"]] == [
        "Electrical",
        "Voltage",
    ]
    assert [u["value"]["raw_text"] for u in contexts[0]["units"]] == ["V"]
    assert "400 V" not in result["text"] and "20 A" not in result["text"]
    assert "fixture:unit-A" not in {s["locator"] for s in result["source_locators"]}
    assert "fixture:unit-V" in {s["locator"] for s in result["source_locators"]}
    assert '"start"' not in result["text"] and '"end"' not in result["text"]
    assert json.dumps(extraction) == before
    assert result == context_for_span(extraction, text, 10, 15)


def test_span_context_matches_text_order_not_table_order(context_table):
    text = "240 V then 400 V"
    result = context_for_span(span_extraction(context_table), text, 0, len(text))
    assert [c["cell"]["raw_text"] for c in json.loads(result["text"])] == ["240 V", "400 V"]


@pytest.mark.parametrize("bounds", [(0, 3), (2, 5), (0, 4)])
def test_span_context_requires_entire_exact_cell(context_table, bounds):
    result = context_for_span(span_extraction(context_table), "400 V", *bounds)
    assert result["status"] == "no_match"
    assert result["text"] == "" and result["source_locators"] == []


def test_span_context_duplicate_value_across_cells_abstains(context_table):
    cells = tuple(
        c.model_copy(update={"raw_text": "400 V"}) if (c.row, c.col) == (3, 1) else c
        for c in context_table.cells
    )
    extraction = span_extraction(context_table.model_copy(update={"cells": cells}))
    result = context_for_span(extraction, "400 V", 0, 5)
    assert result["status"] == "ambiguous"
    assert result["codes"] == ["duplicate_cell_value"]
    assert result["text"] == "" and result["source_locators"] == []


def test_span_context_duplicate_value_in_other_table_abstains(context_table):
    second = context_table.model_copy(
        update={
            "cells": tuple(
                c.model_copy(
                    update={
                        "source": c.source.model_copy(
                            update={"locator": "other:" + c.source.locator}
                        )
                    }
                )
                for c in context_table.cells
            )
        }
    )
    result = context_for_span(span_extraction(context_table, second), "400 V", 0, 5)
    assert result["status"] == "ambiguous" and result["text"] == ""


@pytest.mark.parametrize("bounds", [(0, 5), (8, 13)])
def test_span_context_repeated_text_occurrence_outside_span_is_ambiguous(context_table, bounds):
    result = context_for_span(span_extraction(context_table), "400 V / 400 V", *bounds)
    assert result["status"] == "ambiguous"
    assert result["codes"] == ["repeated_text_match"] and result["text"] == ""


@pytest.mark.parametrize("text", ["1400 V", "400 VA", "400 v", "400  V", "0.4 kV"])
def test_span_context_does_not_match_subtokens_normalize_or_convert(context_table, text):
    result = context_for_span(span_extraction(context_table), text, 0, len(text))
    assert result["status"] == "no_match" and result["text"] == ""


def test_span_context_overlapping_raw_cells_are_ambiguous(context_table):
    cells = tuple(
        c.model_copy(update={"raw_text": "400"}) if (c.row, c.col) == (3, 1) else c
        for c in context_table.cells
    )
    result = context_for_span(
        span_extraction(context_table.model_copy(update={"cells": cells})), "400 V", 0, 5
    )
    assert result["codes"] == ["overlapping_cell_matches"]
    assert result["text"] == ""


def test_span_context_recomputes_header_paths_instead_of_trusting_payload(context_table):
    extraction = span_extraction(context_table)
    raw = next(c for c in extraction["tables"][0]["cells"] if (c["row"], c["col"]) == (2, 1))
    raw["header_paths"][0]["headers"][1]["raw_text"] = "FAKE HEADER"
    result = context_for_span(extraction, "400 V", 0, 5)
    assert result["status"] == "matched"
    assert "FAKE HEADER" not in result["text"] and "Voltage" in result["text"]


@pytest.mark.parametrize("status", ["unknown", "partial"])
def test_span_context_incomplete_inventory_prevents_attribution(context_table, status):
    extraction = span_extraction(context_table)
    extraction["table_structure"]["status"] = status
    result = context_for_span(extraction, "400 V", 0, 5)
    assert result["needs_review"] and result["text"] == ""
    assert result["codes"] == ["incomplete_table_extraction"]


def test_span_context_ocr_cannot_be_cleared_by_summary(context_table):
    extraction = span_extraction(context_table)
    extraction["pages"] = [{"page_no": 2, "method": "ocr", "review_required": False}]
    result = context_for_span(extraction, "400 V", 0, 5)
    assert result["needs_review"] and result["text"] == ""
    assert result["codes"] == ["ocr_review_required"]


def test_span_context_unconfirmed_headers_never_attributed(context_table):
    table = context_table.model_copy(update={"headers_confirmed": False})
    result = context_for_span(span_extraction(table), "400 V", 0, 5)
    assert result["needs_review"] and result["text"] == ""
    assert result["codes"] == ["table_structure_requires_review"]


@pytest.mark.parametrize("bounds", [(-1, 5), (0, 6), (4, 2), (0, 0), (True, 5), (0, 5.0)])
def test_span_context_invalid_offsets_are_not_clamped(context_table, bounds):
    result = context_for_span(span_extraction(context_table), "400 V", *bounds)
    assert result["codes"] == ["invalid_source_span"] and result["text"] == ""


@pytest.mark.parametrize(
    "limits,code",
    [
        (TableLimits(max_text_chars=4), "span_context_text_limit"),
        (TableLimits(max_cells=2), "span_context_cell_limit"),
        (TableLimits(max_output_chars=50), "span_context_metadata_limit"),
        (TableLimits(max_chunk_chars=100), "span_context_output_limit"),
    ],
)
def test_span_context_bounds_fail_without_partial_context(context_table, limits, code):
    result = context_for_span(span_extraction(context_table), "400 V", 0, 5, limits=limits)
    assert result["needs_review"] and result["text"] == "" and result["source_locators"] == []
    assert result["codes"] == [code]


def test_span_context_duplicate_locators_and_malformed_payload_abstain(context_table):
    extraction = span_extraction(context_table)
    extraction["tables"][0]["cells"][1]["source"] = extraction["tables"][0]["cells"][0]["source"]
    assert context_for_span(extraction, "400 V", 0, 5)["codes"] == ["duplicate_cell_locator"]
    extraction["tables"][0]["cells"][0]["row_span"] = -1
    result = context_for_span(extraction, "400 V", 0, 5)
    assert result["codes"] == ["invalid_table_metadata"] and result["text"] == ""


def test_span_context_cyclic_metadata_is_bounded():
    table = {"cells": []}
    table["cycle"] = table
    result = context_for_span({"tables": [table]}, "body", 0, 4)
    assert result["codes"] == ["span_context_metadata_limit"]


def test_span_context_empty_cell_does_not_fill_from_previous_row(context_table):
    cells = tuple(
        c.model_copy(update={"raw_text": "", "row_span": 1}) if c.col == 0 and c.row >= 2 else c
        for c in context_table.cells
    )
    # Supply the actual empty cell instead of a covered merge position.
    cells += (TableCell(row=3, col=0, raw_text="", source=SourceLocator(locator="empty:A4")),)
    table = context_table.model_copy(update={"cells": cells, "header_bands": ((0, 1),)})
    result = context_for_span(span_extraction(table), "240 V", 0, 5)
    assert result["status"] == "matched"
    assert json.loads(result["text"])[0]["row_label"]["raw_text"] == ""


def test_span_context_scan_budget_is_bounded():
    sheet = Workbook().active
    sheet.append(["Item", "Value"])
    for row in range(1, 81):
        sheet.append([f"row_{row}", "a" * row])
    table = table_from_xlsx(sheet, locator="bounded", header_bands=((0,),))
    text = "a" * 20_000
    result = context_for_span(
        span_extraction(table), text, 0, len(text), limits=TableLimits(max_text_chars=20_000)
    )
    assert result["codes"] == ["span_context_scan_limit"]
    assert result["needs_review"] and result["text"] == ""


def test_span_context_match_count_is_bounded():
    sheet = Workbook().active
    sheet.append(["Item", "Value"])
    values = [f"VAL_{row}!" for row in range(33)]
    for row, value in enumerate(values):
        sheet.append([f"row_{row}", value])
    table = table_from_xlsx(sheet, locator="bounded", header_bands=((0,),))
    text = "\n".join(values)
    result = context_for_span(span_extraction(table), text, 0, len(text))
    assert result["codes"] == ["span_context_match_limit"] and result["text"] == ""


def test_indexing_real_table_context_enriches_search_only(tmp_path, monkeypatch):
    import hashlib
    import io

    from api.parsing.attachment_text import extract_xlsx
    from api.specifications import indexing
    from api.specifications.models import Specification
    from api.specifications.validation import sha, validate_spec

    monkeypatch.setenv("WORKSPACE_DATA_DIR", str(tmp_path))
    notice_text = "Device latency 1ms \uc774\ud558"
    notice_path = tmp_path / "notice.txt"
    notice_path.write_text(notice_text, encoding="utf-8")
    origin = {
        "id": "notice",
        "role": "notice_source",
        "text": notice_text,
        "text_origin": "manual_confirmed",
        "storage_path": str(notice_path),
        "sha256": hashlib.sha256(notice_path.read_bytes()).hexdigest(),
    }
    requirement = Specification(
        target="Device",
        property="latency",
        kind="numeric",
        operator="le",
        value="1",
        unit="ms",
        raw=notice_text,
        source={
            "document_id": origin["id"],
            "document_sha256": origin["sha256"],
            "text_sha256": sha(notice_text),
            "start": 0,
            "end": len(notice_text),
            "quote": notice_text,
            "locator": "line 1",
        },
    )
    validation = validate_spec(requirement, origin)
    assert validation["status"] == "valid", validation
    sheet = Workbook().active
    sheet.append(["Item", "Timing"])
    sheet.append([None, "latency (ms)"])
    sheet.append(["Device", "0.5ms"])
    sheet.merge_cells("A1:A2")
    stream = io.BytesIO()
    sheet.parent.save(stream)
    content = stream.getvalue()
    path = tmp_path / "reference.xlsx"
    path.write_bytes(content)
    extracted = extract_xlsx(content)
    doc = {
        "id": "reference",
        "role": "company_reference",
        "storage_path": str(path),
        "sha256": hashlib.sha256(content).hexdigest(),
        "text_origin": "xlsx",
        "text": extracted["text"],
        "extraction": extracted["extraction"],
    }
    inputs = {
        "documents": [origin, doc],
        "specifications": {
            "req": {"payload": requirement.model_dump(mode="json"), "validation": validation}
        },
    }
    rows = indexing.make_chunks(inputs)
    numeric = next(row for row in rows if "0.5ms" in row["quote"])
    assert numeric["payload"]["table_context"]["status"] == "matched"
    assert "Timing" in numeric["payload"]["search_text"]
    assert "latency (ms)" in numeric["payload"]["search_text"]
    assert "Device" in numeric["payload"]["search_text"]
    assert "latency" not in numeric["quote"] and "Device" not in numeric["quote"]
    assert numeric["payload"]["observations"] == [] and numeric["properties"] == []
    assert "".join(row["quote"] for row in rows) == doc["text"]
    for row in rows:
        assert row["quote"] == doc["text"][row["start"] : row["end"]]
        assert row["document_sha256"] == doc["sha256"]
        assert row["text_sha256"] == sha(doc["text"])
        assert row["payload"]["observations"] == []
