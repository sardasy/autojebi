import copy
import io
import json
from datetime import date
from pathlib import Path

import pytest
import sqlalchemy as sa
from openpyxl import Workbook

from alembic.migration import MigrationContext
from alembic.operations import Operations
from alembic.util import load_python_file
from api.workspace import asset_policy, service
from api.workspace.assets import AssetDetails
from api.workspace.tables import company_asset_events, company_assets, workspaces


@pytest.fixture
def asset_env(client, sqlite_engine, tmp_path, monkeypatch):
    from api.config import settings
    monkeypatch.setattr(settings, "upload_dir", str(tmp_path / "uploads"))
    monkeypatch.setenv("WORKSPACE_DATA_DIR", str(tmp_path / "work"))
    monkeypatch.setattr(asset_policy, "today", lambda: date(2026, 9, 11))
    assert client.post("/v2/notices", json={"notice_no":"ASSET-TEST", "title":"Asset test", "source":"TEST"}).status_code == 200
    ws = client.post("/v2/notices/ASSET-TEST/initialize").json()
    ws["inputs"]["company"] = {"business_number":"000-00-00000"}
    with sqlite_engine.begin() as conn:
        conn.execute(sa.update(workspaces).values(inputs=ws["inputs"]))
    return ws


def details(**extra):
    return {"title":"Test company certificate", "category":"business_registration",
            "subject_business_number":"000-00-00000", "issued_on":"2026-09-01",
            "expires_on":"2026-09-20", "no_expiry":False,"active":True, **extra}


def upload(client, **extra):
    out = io.BytesIO()
    workbook = Workbook()
    workbook.active["A1"] = "TEST ONLY"
    workbook.save(out)
    workbook.close()
    result = client.post("/v2/company-assets", data={"metadata":json.dumps(details(**extra))},
                         files={"file":("certificate.xlsx", out.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert result.status_code == 201, result.text
    return result.json()


def link(client, asset, version=1):
    return client.post(f"/v2/notices/ASSET-TEST/company-assets/{asset['id']}",
                       json={"version":version, "asset_version":asset["version"]})


def test_upload_list_alerts_auth_and_hidden_storage(client, asset_env, monkeypatch):
    from api.config import settings
    asset = upload(client)
    assert asset["version"] == 1 and asset["validity"] == "expiring"
    assert asset["details"]["subject_business_number"] == "0000000000"
    assert "storage_path" not in asset["file"]
    result = client.get("/v2/company-assets").json()
    assert result["alerts"]["expiring"] == 1
    assert result["items"][0]["file"]["sha256"] == asset["file"]["sha256"]
    assert client.get(f"/v2/company-assets/{asset['id']}/download").content.startswith(b"PK")
    monkeypatch.setattr(settings,"api_key","test-only-private-key")
    for path in ("/v2/company-assets",f"/v2/company-assets/{asset['id']}/download",f"/v2/company-assets/{asset['id']}/history"):
        assert client.get(path).status_code in {401,403}


def test_link_is_not_approval_and_duplicate_does_not_duplicate_file(client, asset_env):
    asset = upload(client)
    ws = link(client,asset).json()
    assert ws["version"] == 2
    doc = ws["inputs"]["documents"][0]
    assert not doc["confirmed"] and doc["company_asset"]["version"] == 1
    assert doc["role"] == "company_evidence" and doc["text_origin"] == "unread"
    assert link(client,asset,2).json()["version"] == 2
    assert link(client,asset,1).status_code == 409
    assert len(client.get("/v2/notices/ASSET-TEST/workspace").json()["inputs"]["documents"]) == 1


@pytest.mark.parametrize("metadata,code", [
    ({"expires_on":"2026-09-10"},"만료"),
    ({"expires_on":None},"미확인"),
    ({"issued_on":"2026-09-12"},"미도래"),
    ({"subject_business_number":"1111111111"},"불일치"),
])
def test_invalid_company_or_period_blocks_approval_but_allows_review_copy(client,asset_env,metadata,code):
    asset = upload(client,**metadata)
    ws = link(client,asset).json()
    doc = ws["inputs"]["documents"][0]
    text = client.post(f"/v2/notices/ASSET-TEST/documents/{doc['id']}/text",
                       json={"version":2,"text":"TEST certificate text", "note":"Synthetic text"})
    assert text.status_code == 200
    response = client.post(f"/v2/notices/ASSET-TEST/documents/{doc['id']}/review",
        json={"version":3,"confirmed":True,"role":"company_evidence","note":"Test review"})
    assert response.status_code == 409 and code in response.text
    assert client.get(f"/v2/notices/ASSET-TEST/documents/{doc['id']}/download").status_code == 200


def test_metadata_edit_invalidates_link_and_preserves_history(client,asset_env,sqlite_engine):
    asset = upload(client)
    ws = link(client,asset).json()
    inputs = copy.deepcopy(ws["inputs"])
    doc = inputs["documents"][0]
    doc.update(confirmed=True, text="Test body", text_origin="manual_confirmed")
    inputs["proposal_sections"] = [{"id":"s","text":"Test draft", "review_status":"approved"}]
    with sqlite_engine.begin() as conn:
        conn.execute(sa.update(workspaces).values(inputs=inputs))
    updated = client.put(f"/v2/company-assets/{asset['id']}",json={"version":1,"details":details(active=False),"note":"Obsolete test certificate"})
    assert updated.status_code == 200,updated.text
    state = client.get("/v2/notices/ASSET-TEST/workspace").json()
    assert state["version"] == 3
    doc = state["inputs"]["documents"][0]
    assert doc["company_asset"]["stale"] and not doc["confirmed"]
    assert doc["text"] == "Test body"
    assert state["inputs"]["proposal_sections"][0]["review_status"] == "pending"
    history = client.get(f"/v2/company-assets/{asset['id']}/history").json()["items"]
    assert [x["version"] for x in history] == [2,1]
    assert history[-1]["details"]["active"]
    assert client.put(f"/v2/company-assets/{asset['id']}",json={"version":1,"details":details(),"note":"stale"}).status_code == 409
    assert link(client,updated.json(),3).status_code == 409
    assert client.get("/v2/company-assets").json()["items"] == []


def test_reference_cannot_be_relabelled_or_confirmed_as_evidence(client,asset_env):
    asset = upload(client,category="previous_proposal")
    ws = link(client,asset).json()
    assert ws["inputs"]["documents"][0]["role"] == "company_reference"
    assert client.put(f"/v2/company-assets/{asset['id']}",json={"version":1,"details":details(),"note":"Pretend proof"}).status_code == 422
    response = client.post(f"/v2/notices/ASSET-TEST/documents/{ws['inputs']['documents'][0]['id']}/review",
                          json={"version":2,"role":"company_evidence","confirmed":True,"note":"Pretend proof"})
    assert response.status_code == 422


def test_tampered_original_cannot_attach(client,asset_env,sqlite_engine):
    asset = upload(client)
    with sqlite_engine.connect() as conn:
        path = Path(conn.execute(sa.select(company_assets.c.file)).scalar_one()["storage_path"])
    path.write_bytes(b"tampered test fixture")
    assert link(client,asset).status_code == 409
    assert client.get(f"/v2/company-assets/{asset['id']}/download").status_code == 409


def test_invalid_file_or_dates_never_registers_asset(client,asset_env,sqlite_engine):
    result = client.post("/v2/company-assets",data={"metadata":json.dumps(details())},files={"file":("wrong.pdf",b"not pdf")})
    assert result.status_code == 415
    with sqlite_engine.connect() as conn:
        assert conn.execute(sa.select(sa.func.count()).select_from(company_assets)).scalar() == 0
    with pytest.raises(ValueError):
        AssetDetails(**details(expires_on="2026-08-31"))
    with pytest.raises(ValueError):
        AssetDetails(**details(no_expiry=True))


def test_expiry_is_rechecked_without_input_change(monkeypatch):
    doc = {"id":"d", "name":"Certificate", "role":"company_evidence", "text":"Proof", "confirmed":True,
           "item_id":"proof", "company_asset":{**details(),"subject_business_number":"0000000000"}}
    inputs = {"company":{"business_number":"0000000000"},"documents":[doc],"checklist":[{"id":"proof","name":"Proof"}]}
    monkeypatch.setattr(asset_policy,"today",lambda:date(2026,9,20))
    assert not any("필수 증빙" in x for x in service.blockers(inputs))
    monkeypatch.setattr(asset_policy,"today",lambda:date(2026,9,21))
    assert any("만료" in x for x in service.blockers(inputs))
    assert any("필수 증빙" in x for x in service.blockers(inputs))


def test_company_asset_migration_upgrade_and_downgrade():
    migration = load_python_file(str(Path(__file__).resolve().parents[1]/"alembic"/"versions"), "0016_company_assets.py")
    engine = sa.create_engine("sqlite://")
    with engine.begin() as conn:
        with Operations.context(MigrationContext.configure(conn)):
            migration.upgrade()
        for table in (company_assets,company_asset_events):
            assert {c["name"] for c in sa.inspect(conn).get_columns(table.name)} == set(table.c.keys())
        with Operations.context(MigrationContext.configure(conn)):
            migration.downgrade()
        assert "company_assets" not in sa.inspect(conn).get_table_names()
