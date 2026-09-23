import hashlib
import io
import json
import stat
import zipfile
from pathlib import Path

import pytest
from openpyxl import Workbook

from api.parsing.attachment_text import (
    ExtractionError,
    extract_hwpx,
    extract_xlsx,
    extract_zip,
    text_only_table_metadata,
)
from api.specifications.table_structure import TableStructure, build_table_chunks
from api.workspace.source_inventory import source_inventory, source_text_complete


def workbook_bytes(formula=False):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "품목"
    sheet.append(["품명", "규격", "수량", "단가", "금액"])
    sheet.append(["TEST ONLY simulator", "2 kW", 0, 1250.5, "=C2*D2" if formula else 0])
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def archive_bytes(entries):
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w") as archive:
        for name, data in entries:
            archive.writestr(name, data)
    return output.getvalue()


def test_xlsx_preserves_zero_unit_price_and_cell_provenance():
    result = extract_xlsx(workbook_bytes())
    assert result["extraction"]["complete"]
    item = result["extraction"]["items"][0]
    assert item["fields"]["quantity"]["value"] == 0
    assert item["fields"]["unit_price"]["value"] == 1250.5
    assert item["fields"]["unit_price"]["cell"] == "D2"
    assert item["fields"]["specification"]["value"] == "2 kW"
    assert item["review_status"] == "pending"
    assert "시트 품목 · 셀 D2" in result["text"]
    table = result["extraction"]["tables"][0]
    assert table["format"] == "xlsx"
    assert not table["needs_review"]
    quantity = next(c for c in table["cells"] if (c["row"], c["col"]) == (1, 2))
    assert quantity["raw_text"] == "0"
    assert quantity["source"]["locator"].endswith("#sheet=품목!C2")
    assert quantity["header_paths"][0]["headers"][0]["raw_text"] == "수량"


def test_formula_without_cache_is_not_zero_or_evaluated():
    result = extract_xlsx(workbook_bytes(formula=True))
    meta = result["extraction"]
    assert not meta["complete"]
    amount = meta["items"][0]["fields"]["amount"]
    assert amount["value"] is None
    assert amount["formula"] == "=C2*D2"
    assert meta["issues"] == [{"code": "formula_review", "sheet": "품목", "cell": "E2"}]
    assert meta["formula_evaluation"] == "not_performed"
    assert meta["table_structure"]["status"] == "needs_review"
    assert next(c for c in meta["tables"][0]["cells"] if c["source"]["locator"].endswith("!E2"))["raw_text"] == "=C2*D2"


def test_zip_keeps_unread_drawings_in_inventory_not_in_analysis_text():
    excel = workbook_bytes()
    result = extract_zip(archive_bytes([("품목.xlsx", excel), ("도면.dwg", b"inert drawing")]))
    assert not result["extraction"]["complete"]
    members = result["extraction"]["members"]
    assert members[0]["sha256"] == hashlib.sha256(excel).hexdigest()
    assert members[1]["status"] == "unsupported"
    assert "inert drawing" not in result["text"]
    assert "[ZIP 파일 품목.xlsx]" in result["text"]
    doc = {"id": "zip", "role": "notice_source", "text": result["text"],
           "extraction": result["extraction"], "attachment_index": 1, "attachment_url": "url"}
    inputs = {"documents": [doc], "raw": {"ntceSpecDocUrl1": "url", "ntceSpecFileNm1": "archive.zip"}}
    assert source_inventory(inputs)[0]["state"] == "unread"
    assert not source_text_complete(inputs)


@pytest.mark.parametrize("name", ["../escape.xlsx", "/absolute.xlsx", "C:/outside.xlsx", "a\\b.xlsx"])
def test_zip_rejects_unsafe_paths(name):
    entry = zipfile.ZipInfo("placeholder")
    entry.filename = name
    with pytest.raises(ExtractionError, match="unsafe_archive_entry"):
        extract_zip(archive_bytes([(entry, b"x")]))


def test_zip_rejects_symlinks_duplicates_and_keeps_unsupported_nested_members():
    link = zipfile.ZipInfo("linked.xlsx")
    link.external_attr = (stat.S_IFLNK | 0o777) << 16
    with pytest.raises(ExtractionError):
        extract_zip(archive_bytes([(link, b"elsewhere")]))
    with pytest.raises(ExtractionError, match="ambiguous"):
        extract_zip(archive_bytes([("A.xlsx", b"x"), ("a.xlsx", b"y")]))
    result = extract_zip(archive_bytes([("nested.zip", archive_bytes([("a.txt", b"x")]))]))
    assert result["text"] == ""
    assert not result["extraction"]["complete"]
    nested = result["extraction"]["members"][0]["extraction"]["members"][0]
    assert nested["name"] == "a.txt" and nested["status"] == "unsupported"


def test_archive_size_limit_applies_before_member_read(monkeypatch):
    monkeypatch.setattr("api.parsing.attachment_text.MAX_ENTRIES", 1)
    with pytest.raises(ExtractionError, match="archive_size_limit"):
        extract_zip(archive_bytes([("a.dwg", b"x"), ("b.dwg", b"y")]))


def test_spreadsheet_refuses_entities_even_in_utf16():
    with zipfile.ZipFile(io.BytesIO(workbook_bytes())) as original:
        entries = [(e.filename, original.read(e)) for e in original.infolist()]
    entries.append(("customXml/item1.xml", '<!DOCTYPE a [<!ENTITY x "x">]><a>&x;</a>'.encode("utf-16")))
    with pytest.raises(ExtractionError, match="unsafe_spreadsheet_xml"):
        extract_xlsx(archive_bytes(entries))


def test_xlsx_merged_headers_keep_legacy_text_and_source_backed_cells():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Two levels"
    sheet.append(["Item", "Electrical", None])
    sheet.append([None, "Voltage (V)", "Current (A)"])
    sheet.append(["Module", 400, 0])
    sheet.append([None, None, "last"])
    sheet.merge_cells("A1:A2")
    sheet.merge_cells("B1:C1")
    expected = "\n".join(f"[시트 {sheet.title} · 셀 {cell.coordinate}] {cell.value}"
                         for row in sheet for cell in row if cell.value is not None)
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    result = extract_xlsx(output.getvalue())
    assert result["text"] == expected
    table = TableStructure.model_validate(result["extraction"]["tables"][0])
    values = {(c.row, c.col): c for c in table.cells}
    assert values[0, 0].row_span == 2 and values[0, 1].col_span == 2
    assert values[3, 0].raw_text == values[3, 1].raw_text == ""
    assert [h.raw_text for h in values[2, 1].header_paths[0].headers] == ["Electrical", "Voltage (V)"]
    chunks = build_table_chunks(table)
    assert json.loads(chunks[2].text)["row_label"]["raw_text"] == "Module"
    assert not any("start" in c or "end" in c for c in result["extraction"]["tables"][0]["cells"])


def test_xlsx_large_merge_is_bounded_without_losing_old_text():
    with zipfile.ZipFile(io.BytesIO(workbook_bytes())) as archive:
        entries = [(e.filename, archive.read(e)) for e in archive.infolist()]
    entries = [(name, data.replace(b'</worksheet>', b'<mergeCells><mergeCell ref="A1:XFD1048576"/></mergeCells></worksheet>')
                if name == "xl/worksheets/sheet1.xml" else data) for name, data in entries]
    result = extract_xlsx(archive_bytes(entries))
    assert "TEST ONLY simulator" in result["text"]
    assert result["extraction"]["complete"]
    assert result["extraction"]["tables"] == []
    assert result["extraction"]["table_structure"]["issues"][0]["code"] == "table_dimension_limit"


def test_zip_exposes_tables_at_top_level_with_distinct_member_sources():
    content = workbook_bytes()
    result = extract_zip(archive_bytes([("first.xlsx", content), ("second.xlsx", content)]))
    tables = result["extraction"]["tables"]
    assert len(tables) == 2
    assert tables[0]["source"]["locator"] != tables[1]["source"]["locator"]
    assert "/first.xlsx#" in tables[0]["source"]["locator"]
    assert tables[0] == result["extraction"]["members"][0]["extraction"]["tables"][0]


def test_nested_zip_reads_hwpx_and_keeps_container_chain_and_hashes():
    hwpx = hwpx_table_bytes()
    nested = archive_bytes([("spec.hwpx", hwpx)])
    result = extract_zip(archive_bytes([("bundle.zip", nested), ("spec.hwpx", hwpx)]))
    assert result["extraction"]["complete"]
    members = result["extraction"]["members"]
    assert members[0]["sha256"] == hashlib.sha256(nested).hexdigest()
    assert members[0]["extraction"]["members"][0]["sha256"] == hashlib.sha256(hwpx).hexdigest()
    assert result["text"].count("Specification Voltage 400 V") == 2
    assert "[ZIP 파일 bundle.zip]\n[ZIP 파일 spec.hwpx]" in result["text"]
    tables = result["extraction"]["tables"]
    assert len(tables) == 2 and tables[0]["source"] != tables[1]["source"]
    assert "/bundle.zip/zip/spec.hwpx/Contents/section0.xml" in tables[0]["source"]["locator"]
    assert tables[0] == members[0]["extraction"]["members"][0]["extraction"]["tables"][0]


def test_nested_archive_depth_limit_keeps_other_sources_and_blocks_completeness():
    deepest = archive_bytes([("spec.hwpx", hwpx_table_bytes())])
    for _ in range(3):
        deepest = archive_bytes([("nested.zip", deepest)])
    result = extract_zip(archive_bytes([("too-deep.zip", deepest), ("readable.hwpx", hwpx_table_bytes())]))
    assert "Specification Voltage 400 V" in result["text"]
    assert not result["extraction"]["complete"]
    member = result["extraction"]["members"][0]
    for _ in range(2):
        member = member["extraction"]["members"][0]
    assert member["status"] == "blocked" and member["error_code"] == "archive_depth_limit"


def test_nested_archive_rechecks_unsafe_paths_before_member_processing():
    content = archive_bytes([("nested.zip", archive_bytes([("../escape.hwpx", hwpx_table_bytes())]))])
    with pytest.raises(ExtractionError, match="unsafe_archive_entry"):
        extract_zip(content)


def test_archive_office_members_share_expansion_and_entry_limits(monkeypatch):
    content = archive_bytes([("first.hwpx", hwpx_table_bytes()), ("second.hwpx", hwpx_table_bytes())])
    monkeypatch.setattr("api.parsing.attachment_text.MAX_BYTES", len(content) + 10)
    with pytest.raises(ExtractionError, match="archive_size_limit"):
        extract_zip(content)
    monkeypatch.setattr("api.parsing.attachment_text.MAX_BYTES", 50_000_000)
    monkeypatch.setattr("api.parsing.attachment_text.MAX_ENTRIES", 3)
    with pytest.raises(ExtractionError, match="archive_entry_limit"):
        extract_zip(content)


def test_nested_zip_shares_pdf_ocr_budget(monkeypatch):
    budgets = []
    def pdf(content, *, ocr_budget, locator):
        budgets.append(ocr_budget[0])
        ocr_budget[0] -= 1
        return {"text": "recognized", "extraction": {"complete": False,
            **text_only_table_metadata(locator=locator)}}
    monkeypatch.setattr("api.parsing.pdf_text.extract_pdf", pdf)
    result = extract_zip(archive_bytes([("a.pdf", b"%PDF first"),
        ("nested.zip", archive_bytes([("b.pdf", b"%PDF second")]))]), ocr_budget=[2])
    assert budgets == [2, 1] and not result["extraction"]["complete"]


def test_zip_hwp_requires_worker_and_shares_call_budget_across_nesting():
    hwp = bytes.fromhex("d0cf11e0a1b11ae1") + b"inert fixture"
    content = archive_bytes([("a.hwp", hwp), ("nested.zip", archive_bytes([("b.hwp", hwp)]))])
    missing = extract_zip(content)
    assert missing["extraction"]["members"][0]["error_code"] == "hwp_worker_required"
    seen = []
    def worker(data, locator):
        seen.append((data, locator))
        return {"text": "HWP body", "extraction": {"version": 1, "complete": True,
            **text_only_table_metadata(locator=locator)}}
    result = extract_zip(content, hwp_extractor=worker, hwp_budget=[1])
    assert len(seen) == 1 and seen[0][0] == hwp
    assert "HWP body" in result["text"] and not result["extraction"]["complete"]
    nested = result["extraction"]["members"][1]["extraction"]["members"][0]
    assert nested["error_code"] == "hwp_member_limit"


def test_zip_matching_extension_requires_matching_content():
    result = extract_zip(archive_bytes([("document.hwpx", b"not an archive")]))
    assert result["extraction"]["members"][0]["error_code"] == "member_format_mismatch"
    assert result["text"] == "" and not result["extraction"]["complete"]


def hwpx_table_bytes(span=1):
    xml = (
        '<hp:sec xmlns:hp="http://www.hancom.co.kr/hwpml/2011/paragraph"><hp:p><hp:run>'
        '<hp:t>Specification</hp:t><hp:tbl rowCnt="2" colCnt="1"><hp:tr><hp:tc>'
        '<hp:subList><hp:p><hp:run><hp:t>Voltage</hp:t></hp:run></hp:p></hp:subList>'
        f'<hp:cellAddr rowAddr="0" colAddr="0"/><hp:cellSpan rowSpan="1" colSpan="{span}"/>'
        '</hp:tc></hp:tr><hp:tr><hp:tc><hp:subList><hp:p><hp:run><hp:t>400 V</hp:t>'
        '</hp:run></hp:p></hp:subList><hp:cellAddr rowAddr="1" colAddr="0"/>'
        '<hp:cellSpan rowSpan="1" colSpan="1"/></hp:tc></hp:tr></hp:tbl></hp:run></hp:p></hp:sec>'
    )
    return archive_bytes([("Contents/section0.xml", xml)])


def test_hwpx_real_sections_preserve_text_order_and_add_actual_tables():
    result = extract_hwpx(hwpx_table_bytes())
    assert result["text"] == "Specification Voltage 400 V"
    assert result["extraction"]["complete"]
    table = result["extraction"]["tables"][0]
    assert table["needs_review"]  # Header role is not asserted by text order.
    assert table["cells"][1]["raw_text"] == "400 V"
    assert "/Contents/section0.xml#table=0/cell=1;r=1;c=0" in table["cells"][1]["source"]["locator"]
    assert table["cells"][1]["source"]["page"] is None
    assert TableStructure.model_validate(table).format == "hwpx"


def test_hwpx_structure_error_keeps_body_and_reports_rejected_source():
    result = extract_hwpx(hwpx_table_bytes(span=2))
    assert result["text"] == "Specification Voltage 400 V"
    assert result["extraction"]["tables"] == []
    meta = result["extraction"]["table_structure"]
    assert meta["detected_count"] == 1 and meta["needs_review"]
    assert meta["issues"][0]["code"] == "span_out_of_bounds"
    assert "#table=0" in meta["issues"][0]["source"]["locator"]


def test_hwpx_plain_prose_is_not_a_table_review_failure():
    result = extract_hwpx(archive_bytes([("Contents/section0.xml", "<section><text>Plain prose</text></section>")]))
    assert result["text"] == "Plain prose"
    assert result["extraction"]["table_structure"]["status"] == "not_detected"
    assert not result["extraction"]["table_structure"]["needs_review"]


def test_hwpx_execution_uses_secure_parser():
    content = archive_bytes([("Contents/section0.xml", '<!DOCTYPE a [<!ENTITY x "unsafe">]><a>&x;</a>')])
    from defusedxml.common import DefusedXmlException

    with pytest.raises(DefusedXmlException):
        extract_hwpx(content)


@pytest.mark.parametrize("count,review,status", [(None, False, "unknown"), (0, False, "unknown"), (2, True, "needs_review")])
def test_hwp_text_does_not_assert_tables_without_positive_evidence(count, review, status):
    metadata = text_only_table_metadata(locator="hwp:fixture", table_count=count)
    assert metadata["tables"] == []
    assert metadata["table_structure"]["status"] == status
    assert metadata["table_structure"]["needs_review"] is review
    assert metadata["table_structure"]["structure_trustworthy"] is False


def test_metadata_budget_failure_is_explicit_and_preserves_legacy_text(monkeypatch):
    monkeypatch.setattr("api.parsing.attachment_text.MAX_TABLE_METADATA", 10)
    result = extract_xlsx(workbook_bytes())
    assert "TEST ONLY simulator" in result["text"]
    assert result["extraction"]["complete"]
    assert result["extraction"]["table_structure"]["issues"][0]["code"] == "table_metadata_limit"


@pytest.mark.parametrize("cached", [False, True])
def test_execution_hwpx_adds_structure_without_rewriting_cached_body(tmp_path, monkeypatch, cached):
    from api.workspace.execution import extract

    monkeypatch.setenv("WORKSPACE_DATA_DIR", str(tmp_path))
    path = tmp_path / "source.hwpx"
    content = hwpx_table_bytes()
    path.write_bytes(content)
    doc = {"id": "source", "name": path.name, "storage_path": str(path),
           "sha256": hashlib.sha256(content).hexdigest(), "extraction": {"old_metadata": "keep"}}
    if cached:
        doc.update(text="Earlier exact body", text_origin="hwpx")
        doc["extraction"].update(text_rule_version=2, complete=True)
    result = extract({"snapshot": {"documents": [doc]}})["documents"][0]
    assert result["text"] == ("Earlier exact body" if cached else "Specification Voltage 400 V")
    assert result["extraction"]["old_metadata"] == "keep"
    assert len(result["extraction"]["tables"]) == 1
    assert result["extraction"]["tables"][0]["cells"][1]["raw_text"] == "400 V"


def test_execution_cached_hwp_retains_metadata_and_does_not_reopen_worker(tmp_path, monkeypatch):
    from api.workspace import execution

    monkeypatch.setenv("WORKSPACE_DATA_DIR", str(tmp_path))
    path = tmp_path / "source.hwp"
    path.write_bytes(bytes.fromhex("d0cf11e0a1b11ae1") + b"inert")
    doc = {"id": "source", "name": path.name, "storage_path": str(path),
           "sha256": hashlib.sha256(path.read_bytes()).hexdigest(), "text_origin": "hwp_ole",
           "text": "Cached body", "extraction": {"custom": ["keep"]}}
    monkeypatch.setattr("api.workspace.executors.extract.HwpAgentClient", lambda **_: pytest.fail("cached HWP must not reopen COM"))
    result = execution.extract({"snapshot": {"documents": [doc]}})["documents"][0]
    assert result["text"] == "Cached body"
    assert result["extraction"]["custom"] == ["keep"]
    assert result["extraction"]["table_structure"]["status"] == "unknown"
    assert not result["extraction"]["table_structure"]["needs_review"]


def test_execution_zip_hwp_uses_authenticated_worker_with_immutable_digest_source(tmp_path, monkeypatch):
    from api.workspace import execution

    monkeypatch.setenv("WORKSPACE_DATA_DIR", str(tmp_path))
    monkeypatch.setenv("HWP_AGENT_TOKEN", "test-only-token")
    hwp = bytes.fromhex("d0cf11e0a1b11ae1") + b"inert fixture"
    content = archive_bytes([("folder/Korean source.hwp", hwp)])
    path = tmp_path / "sources.zip"
    path.write_bytes(content)
    opened = []
    class Worker:
        def analyze_document(self, source):
            source = Path(source)
            assert source.parent == tmp_path / "uploads" / "archive-members"
            assert source.name == hashlib.sha256(hwp).hexdigest() + ".hwp"
            assert source.read_bytes() == hwp
            opened.append(source)
            return {"text": "Actual Worker body", "table_count": 1}
    monkeypatch.setattr("api.workspace.executors.extract.HwpAgentClient", lambda **_: Worker())
    doc = {"id": "zip", "name": path.name, "storage_path": str(path),
           "sha256": hashlib.sha256(content).hexdigest()}
    result = execution.extract({"snapshot": {"documents": [doc]}})["documents"][0]
    assert "Actual Worker body" in result["text"] and result["extraction"]["complete"]
    assert opened[0].is_file()  # Preserve the source even if the Worker later needs recovery.
    member = result["extraction"]["members"][0]
    assert member["sha256"] == hashlib.sha256(hwp).hexdigest()
    assert member["extraction"]["table_structure"]["status"] == "needs_review"
    assert not member["extraction"]["table_structure"]["structure_trustworthy"]
    assert "%2F" in member["extraction"]["table_structure"]["issues"][0]["source"]["locator"]


def test_archive_hwp_refuses_changed_retained_source_without_calling_worker(tmp_path, monkeypatch):
    from api.parsing.attachment_text import ArchiveMemberError
    from api.workspace import execution

    monkeypatch.setenv("WORKSPACE_DATA_DIR", str(tmp_path))
    monkeypatch.setenv("HWP_AGENT_TOKEN", "test-only-token")
    content = bytes.fromhex("d0cf11e0a1b11ae1") + b"original"
    folder = tmp_path / "uploads" / "archive-members"
    folder.mkdir(parents=True)
    (folder / (hashlib.sha256(content).hexdigest() + ".hwp")).write_bytes(b"changed")
    monkeypatch.setattr("api.workspace.executors.extract.HwpAgentClient", lambda **_: pytest.fail("changed source"))
    with pytest.raises(ArchiveMemberError, match="hwp_member_source_changed"):
        execution.extract_archive_hwp(content, "test:member")


def test_archive_hwp_missing_auth_is_an_explicit_incomplete_member(tmp_path, monkeypatch):
    from api.workspace.execution import extract_archive_hwp

    monkeypatch.setenv("WORKSPACE_DATA_DIR", str(tmp_path))
    monkeypatch.delenv("HWP_AGENT_TOKEN", raising=False)
    hwp = bytes.fromhex("d0cf11e0a1b11ae1") + b"inert fixture"
    result = extract_zip(archive_bytes([("source.hwp", hwp)]), hwp_extractor=extract_archive_hwp)
    assert not result["extraction"]["complete"]
    assert result["extraction"]["members"][0]["error_code"] == "hwp_worker_auth_missing"
    assert not (tmp_path / "uploads").exists()


def test_execution_cached_pdf_only_enriches_structure_keys(tmp_path, monkeypatch):
    from api.parsing import pdf_text
    from api.workspace import execution

    monkeypatch.setenv("WORKSPACE_DATA_DIR", str(tmp_path))
    path = tmp_path / "source.pdf"
    path.write_bytes(b"%PDF fixture")
    previous = {"version": 1, "complete": True, "pages": [{"old": "page"}], "custom": "keep"}
    doc = {"id": "source", "name": path.name, "storage_path": str(path),
           "sha256": hashlib.sha256(path.read_bytes()).hexdigest(), "text_origin": "pdf",
           "text": "Exact cached body", "extraction": previous}
    monkeypatch.setattr(pdf_text, "extract_pdf", lambda *_, **__: {
        "text": "New extraction ordering", "extraction": {"version": 1, "complete": False,
        "tables": [], "table_structure": {"version": 1, "status": "unknown"}}})
    result = execution.extract({"snapshot": {"documents": [doc]}})["documents"][0]
    assert result["text"] == "Exact cached body"
    assert all(result["extraction"][key] == value for key, value in previous.items())
    assert result["extraction"]["table_structure"]["status"] == "unknown"
    assert doc["extraction"] == previous
