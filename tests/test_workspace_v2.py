from datetime import UTC, datetime, timedelta

import pytest
from sqlalchemy import select, update

from api.routers.notices import bid_pipeline
from api.workspace import service
from api.workspace.errors import WorkspaceError
from api.workspace.tables import jobs, workspaces
from api.workspace.worker import run_once


@pytest.fixture
def workspace(client, tmp_path, monkeypatch):
    monkeypatch.setenv("WORKSPACE_DATA_DIR", str(tmp_path / "workspace"))
    from api.config import settings

    monkeypatch.setattr(settings, "upload_dir", str(tmp_path / "uploads"))
    monkeypatch.setattr(settings, "workspace_llm_provider", "anthropic")
    monkeypatch.setattr(settings, "google_api_key", "")
    monkeypatch.setattr(settings, "anthropic_api_key", "")
    response = client.post(
        "/v2/notices",
        json={
            "notice_no": "E2E-V2",
            "title": "Source test",
            "source": "MANUAL_E2E",
            "raw": {"ntceInsttNm": "기관", "bidClseDt": "2026-09-23 10:00:00", "presmptPrce": 999},
        },
    )
    assert response.status_code == 200, response.text
    response = client.post("/v2/notices/E2E-V2/initialize")
    assert response.status_code == 200, response.text
    return response.json()


def enqueue(client, kind="excel", version=1, key="one"):
    return client.post(
        "/v2/notices/E2E-V2/jobs",
        json={"kind": kind, "input_version": version, "idempotency_key": key},
    )


@pytest.mark.parametrize("value,blocked", [
    (None, True), ("invalid", True), ("2026-09-14T10:00:00", True),
    ("2026-09-14T10:00:00+09:00", True),
    ("2026-09-14T10:00:01+09:00", False),
])
def test_deadline_is_timezone_aware_and_boundary_is_closed(monkeypatch, value, blocked):
    monkeypatch.setattr(service, "now", lambda: datetime(2026, 9, 14, 1, tzinfo=UTC))
    assert bool(service.deadline_issues({"notice": {"close_date": value}})) is blocked


def test_expired_notice_cannot_be_ready_but_keeps_review_download(client, workspace, sqlite_engine):
    record = enqueue(client).json()
    run_once(sqlite_engine)
    with sqlite_engine.begin() as conn:
        inputs = workspace["inputs"]
        inputs["notice"]["close_date"] = "2000-01-01T00:00:00+09:00"
        conn.execute(update(workspaces).values(inputs=inputs))
    response = client.post("/v2/notices/E2E-V2/reviews", json={
        "version": 1, "job_id": record["id"], "action": "ready", "note": "Check deadline",
    })
    assert response.status_code == 409
    assert response.json()["detail"]["code"] == "deadline_unavailable"
    ws = client.get("/v2/notices/E2E-V2/workspace").json()
    assert not ws["ready"] and "마감 경과" in ws["blockers"][0]
    assert client.get(f"/v2/jobs/{record['id']}/download").status_code == 200


@pytest.fixture
def reference_workspace(workspace, sqlite_engine):
    inputs = workspace["inputs"]
    inputs["documents"] = [
        {"id": "old", "name": "Previous proposal", "role": "company_reference",
         "text": "Previous delivery used three units.", "confirmed": False},
        {"id": "source", "name": "Notice", "role": "notice_source", "text": "Need three units."},
    ]
    inputs["requirements"] = [{"id": "req", "text": "Need three units.", "document_id": "source",
                               "quote": "Need three units.", "locator": "p1", "mandatory": True,
                               "review_status": "pending"}]
    with sqlite_engine.begin() as conn:
        conn.execute(update(workspaces).values(inputs=inputs))
    return {"version": 1, "requirement_id": "req", "document_id": "old",
            "quote": "Previous delivery used three units.", "text": "We propose three units.",
            "note": "Adapted for current notice; proof still needed"}


def test_reference_draft_keeps_provenance_and_cannot_approve(client, reference_workspace):
    result = client.post("/v2/notices/E2E-V2/reference-drafts", json=reference_workspace)
    assert result.status_code == 200, result.text
    state = result.json()
    section = state["inputs"]["proposal_sections"][0]
    assert state["version"] == 2 and not state["ready"]
    assert section["review_status"] == "pending" and section["evidence_ids"] == []
    assert section["reference"]["quote"] == reference_workspace["quote"]
    assert section["reference"]["input_version"] == 1
    review = client.post(f"/v2/notices/E2E-V2/sections/{section['id']}/review", json={
        "version": 2, "text": section["text"], "approved": True, "note": "Checked",
        "evidence_ids": ["old"], "requirement_ids": ["req"],
    })
    assert review.status_code == 409
    assert client.post("/v2/notices/E2E-V2/reference-drafts", json=reference_workspace).status_code == 409
    assert len(client.get("/v2/notices/E2E-V2/workspace").json()["inputs"]["proposal_sections"]) == 1


@pytest.mark.parametrize("field,value", [
    ("quote", "Invented quote"), ("quote", " "), ("text", " "),
    ("document_id", "source"), ("document_id", "another-notice"),
    ("requirement_id", "another-notice"),
])
def test_reference_draft_rejects_invalid_inputs(client, reference_workspace, field, value):
    body = {**reference_workspace, field: value}
    assert client.post("/v2/notices/E2E-V2/reference-drafts", json=body).status_code == 422
    assert client.get("/v2/notices/E2E-V2/workspace").json()["version"] == 1


def test_previous_writing_cannot_be_reclassified_as_evidence(client, reference_workspace):
    result = client.post("/v2/notices/E2E-V2/documents/old/review", json={
        "version": 1, "role": "company_evidence", "confirmed": True, "note": "Checked",
    })
    assert result.status_code == 422
    assert result.json()["detail"]["code"] == "reference_not_evidence"


def test_section_draft_save_is_pending_not_rejected(client, reference_workspace):
    state = client.post("/v2/notices/E2E-V2/reference-drafts", json=reference_workspace).json()
    section = state["inputs"]["proposal_sections"][0]
    response = client.post(f"/v2/notices/E2E-V2/sections/{section['id']}/review", json={
        "version": state["version"], "text": "Edited draft", "note": "Work in progress", "approved": None,
    })
    assert response.status_code == 200
    assert response.json()["inputs"]["proposal_sections"][0]["review_status"] == "pending"


def test_pdf_preview_is_inline_but_download_remains_attachment(client, workspace):
    uploaded = client.post("/v2/notices/E2E-V2/documents", data={"version": "1", "role": "notice_source"}, files={"file": ("sample.pdf", b"%PDF-1.4\n%%EOF", "application/pdf")})
    assert uploaded.status_code == 200, uploaded.text
    doc = uploaded.json()["inputs"]["documents"][0]
    path = f"/v2/notices/E2E-V2/documents/{doc['id']}/download"
    assert client.get(path).headers["content-disposition"].startswith("attachment;")
    preview = client.get(path + "?preview=true")
    assert preview.headers["content-disposition"].startswith("inline;")
    assert preview.headers["content-type"] == "application/pdf"


def test_excel_extraction_metadata_survives_job_apply(client, workspace, sqlite_engine):
    import io

    from openpyxl import Workbook

    book = Workbook()
    book.active.append(["품명", "수량", "금액"])
    book.active.append(["TEST ONLY", 2, "=B2*100"])
    data = io.BytesIO()
    book.save(data)
    book.close()
    uploaded = client.post("/v2/notices/E2E-V2/documents", data={"version": "1", "role": "notice_source"},
                           files={"file": ("items.xlsx", data.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert uploaded.status_code == 200
    version = uploaded.json()["version"]
    queued = enqueue(client, kind="extract", version=version).json()
    assert run_once(sqlite_engine)
    applied = client.post(f"/v2/notices/E2E-V2/jobs/{queued['id']}/apply", json={"version": version})
    assert applied.status_code == 200, applied.text
    inputs = applied.json()["inputs"]
    doc = inputs["documents"][0]
    assert doc["text_origin"] == "xlsx"
    assert doc["extraction"]["complete"] is False
    assert doc["extraction"]["items"][0]["fields"]["amount"]["value"] is None
    assert doc["extraction"]["tables"][0]["format"] == "xlsx"
    assert doc["extraction"]["table_structure"]["status"] == "needs_review"
    formula = next(cell for cell in doc["extraction"]["tables"][0]["cells"]
                   if (cell["row"], cell["col"]) == (1, 2))
    assert formula["raw_text"] == "=B2*100"
    assert formula["source"]["locator"].endswith("!C2")
    assert "TEST ONLY" in doc["text"]
    assert doc["confirmed"] is False
    assert not inputs["source_reviewed"]


def test_pdf_page_preview_renders_and_checks_page_bounds(client, workspace):
    from pathlib import Path

    pdf = Path(__file__).resolve().parents[1] / "frontend/e2e/fixtures/technical-spec-sample.pdf"
    uploaded = client.post("/v2/notices/E2E-V2/documents", data={"version": "1", "role": "notice_source"}, files={"file": ("source.pdf", pdf.read_bytes(), "application/pdf")}).json()
    doc = uploaded["inputs"]["documents"][0]
    path = f"/v2/notices/E2E-V2/documents/{doc['id']}/preview"
    result = client.get(path)
    assert result.status_code == 200
    assert result.headers["content-type"] == "image/png"
    assert result.content.startswith(b"\x89PNG")
    assert int(result.headers["x-page-count"]) >= 1
    assert client.get(path + "?page=999").status_code == 404
    assert client.get(path + "?page=0").status_code == 422


def test_normalized_facts_not_estimated_price(workspace):
    assert workspace["inputs"]["notice"]["org_name"] == "기관"
    assert workspace["inputs"]["notice"]["close_date"].startswith("2026-09-23")
    assert workspace["inputs"]["notice"]["base_price"] is None
    assert workspace["decision"] == "pending"
    assert not workspace["ready"]


def test_idempotency_and_version_conflict(client, workspace):
    first = enqueue(client).json()
    assert enqueue(client).json()["id"] == first["id"]
    assert enqueue(client, key="second-click").json()["id"] == first["id"]
    assert enqueue(client, kind="hwp").status_code == 409
    assert enqueue(client, version=999, key="stale").status_code == 409


def test_failed_llm_does_not_change_decision_or_legacy_score(client, workspace, sqlite_engine):
    job = enqueue(client, kind="analyze").json()
    assert run_once(sqlite_engine)
    result = client.get(f"/v2/jobs/{job['id']}").json()
    assert result["status"] == "blocked"
    ws = client.get("/v2/notices/E2E-V2/workspace").json()
    assert ws["decision"] == "pending"
    with sqlite_engine.connect() as conn:
        row = (
            conn.execute(select(bid_pipeline).where(bid_pipeline.c.notice_no == "E2E-V2"))
            .mappings()
            .one()
        )
        assert row["status"] == "collected"
        assert row["analysis"] == {}


def test_dry_run_preserves_all_business_state(client, workspace, sqlite_engine):
    job = enqueue(client, kind="notify").json()
    run_once(sqlite_engine)
    done = client.get(f"/v2/jobs/{job['id']}").json()
    assert done["result"]["dry_run"] is True
    assert done["result"]["delivered"] is False
    ws = client.get("/v2/notices/E2E-V2/workspace").json()
    assert ws["version"] == 1 and ws["decision"] == "pending"


def test_late_result_cannot_become_current_success(client, workspace, sqlite_engine, monkeypatch):
    def execute(engine, job):
        with engine.begin() as conn:
            conn.execute(update(workspaces).values(version=2))
        return {"warnings": []}

    monkeypatch.setattr("api.workspace.worker.execute", execute)
    job = enqueue(client).json()
    run_once(sqlite_engine)
    done = client.get(f"/v2/jobs/{job['id']}").json()
    assert done["status"] == "blocked"
    assert done["error"]["code"] == "version_conflict"


def test_expired_worker_requires_manual_recovery(client, workspace, sqlite_engine):
    job = enqueue(client, kind="hwp").json()
    with sqlite_engine.begin() as conn:
        conn.execute(
            update(jobs)
            .where(jobs.c.id == job["id"])
            .values(status="running", lease_until=service.now() - timedelta(minutes=1))
        )
    assert not run_once(sqlite_engine)
    result = client.get(f"/v2/jobs/{job['id']}").json()
    assert result["status"] == "failed"
    assert result["error"]["code"] == "interrupted"


def test_excel_download_review_and_stale_approval(client, workspace, sqlite_engine):
    job = enqueue(client).json()
    run_once(sqlite_engine)
    result = client.get(f"/v2/jobs/{job['id']}").json()
    assert result["status"] == "succeeded" and result["quality"] == "warning"
    assert client.get(f"/v2/jobs/{job['id']}/download").content[:2] == b"PK"
    review = {"version": 1, "job_id": job["id"], "action": "approved", "note": "검토 초안 확인"}
    assert client.post("/v2/notices/E2E-V2/reviews", json=review).status_code == 200
    assert (
        client.post("/v2/notices/E2E-V2/reviews", json={**review, "action": "ready"}).status_code
        == 409
    )
    assert (
        client.post(
            "/v2/notices/E2E-V2/decision",
            json={"version": 1, "decision": "hold", "note": "확인 필요"},
        ).status_code
        == 200
    )
    assert (
        client.post("/v2/notices/E2E-V2/reviews", json={**review, "version": 2}).status_code == 409
    )
    assert (
        client.post(
            "/v2/notices/E2E-V2/decision",
            json={"version": 1, "decision": "decline", "note": "stale"},
        ).status_code
        == 409
    )


def test_public_notice_never_company_evidence(client, workspace):
    result = client.post(
        "/v2/notices/E2E-V2/documents",
        data={"version": 1, "role": "notice_source"},
        files={"file": ("notice.pdf", b"%PDF-test", "application/pdf")},
    )
    assert result.status_code == 200, result.text
    doc = result.json()["inputs"]["documents"][0]
    assert doc["confirmed"] is False
    result = client.post(
        f"/v2/notices/E2E-V2/documents/{doc['id']}/review",
        json={
            "version": 2,
            "role": "company_evidence",
            "confirmed": True,
            "note": "사업자 단어 발견",
        },
    )
    assert result.status_code == 409
    assert result.json()["detail"]["code"] == "source_not_evidence"


def test_quote_mismatch_rejected(client, workspace):
    response = client.post(
        "/v2/notices/E2E-V2/requirements",
        json={
            "version": 1,
            "text": "invented",
            "document_id": "missing",
            "quote": "fake",
            "locator": "page1",
            "category": "technical",
        },
    )
    assert response.status_code == 409


def test_upstream_change_blocks_jobs_and_preserves_history(client, workspace, sqlite_engine):
    with sqlite_engine.begin() as conn:
        conn.execute(update(bid_pipeline).values(title="Changed"))
    response = enqueue(client)
    assert response.status_code == 409
    assert client.get("/v2/notices/E2E-V2/workspace").json()["source_changed"] is True
    result = client.post("/v2/notices/E2E-V2/sync", json={"version": 1})
    assert result.json()["version"] == 2
    assert len(result.json()["reviews"]) == 2


def test_failed_hwp_visible_and_not_approvable(client, workspace, sqlite_engine):
    job = enqueue(client, kind="hwp").json()
    run_once(sqlite_engine)
    result = client.get(f"/v2/jobs/{job['id']}").json()
    assert result["status"] == "blocked"
    response = client.post(
        "/v2/notices/E2E-V2/reviews",
        json={"version": 1, "job_id": job["id"], "action": "approved", "note": "cannot"},
    )
    assert response.status_code == 409


def test_llm_budget_is_durable_and_bounded(client, workspace, sqlite_engine, monkeypatch):

    from api.config import settings
    from api.workspace.execution import reserve_call

    monkeypatch.setattr(settings, "anthropic_api_key", "test-no-network")
    job = {"notice_no": "E2E-V2", "input_version": 1}
    reserve_call(sqlite_engine, job)
    reserve_call(sqlite_engine, job)
    with pytest.raises(WorkspaceError) as error:
        reserve_call(sqlite_engine, job)
    assert error.value.detail["code"] == "llm_budget_exhausted"


def test_legacy_analysis_failure_preserves_documents(client, workspace, sqlite_engine, monkeypatch):
    from api.services.claude_analyzer import AnalyzeResult

    with sqlite_engine.begin() as conn:
        conn.execute(
            update(bid_pipeline).values(
                fit_score=85,
                status="documents_analyzed",
                analysis={"document_automation": {"uploads": [{"id": "original"}]}},
            )
        )

    class Failure:
        def analyze_notice(self, **kwargs):
            return AnalyzeResult(category="비관련", fit_score=0, analysis={"error": "no auth"})

    monkeypatch.setattr("api.routers.notices.ClaudeAnalyzer", Failure)
    assert client.post("/notices/E2E-V2/analyze").status_code == 502
    with sqlite_engine.connect() as conn:
        row = conn.execute(select(bid_pipeline)).mappings().one()
    assert row["fit_score"] == 85 and row["status"] == "documents_analyzed"
    assert row["analysis"]["document_automation"]["uploads"][0]["id"] == "original"


def test_reanalysis_never_unblocks_rejected_upload():
    from api.models.notices import DocumentChecklistItem
    from api.services.document_automation import _apply_uploads_to_checklist

    item = DocumentChecklistItem(
        id="biz", name="사업자등록증", type="company_common", status="blocked", source="manual"
    )
    result = _apply_uploads_to_checklist(
        [item], [{"id": "public", "item_id": "biz", "source_ref": "g2b_attachment"}]
    )
    assert result[0].status == "blocked"


def test_statement_of_storage_is_not_spec_evidence():
    from api.services.spec_items import build_spec_item_candidates

    row = {
        "raw": {},
        "analysis": {
            "document_automation": {
                "uploads": [
                    {
                        "name": "notice.hwp",
                        "analysis_summary": "파일을 저장했습니다",
                        "text_extract_error": "offline",
                    }
                ]
            }
        },
    }
    candidates = build_spec_item_candidates(row)
    assert not any(c["required_value"] for c in candidates)
    assert not any("파일을 저장" in str(c["source_text"]) for c in candidates)


def test_hwp_mock_writes_audited_artifact_and_remaining_fields(
    client, workspace, sqlite_engine, tmp_path, monkeypatch
):
    import zipfile

    from api.routers.notices import document_field_mappings, hwp_templates
    from api.services.hwp_agent_client import PutFieldsOutcome

    def write_hwpx(path):
        with zipfile.ZipFile(path, "w") as archive:
            archive.writestr("mimetype", "application/hwp+zip")
            archive.writestr(
                "Contents/header.xml",
                '<hh:head xmlns:hh="http://www.hancom.co.kr/hwpml/2011/head" secCnt="1"/>',
            )
            archive.writestr(
                "Contents/section0.xml",
                '<hs:sec xmlns:hs="http://www.hancom.co.kr/hwpml/2011/section" '
                'xmlns:hp="http://www.hancom.co.kr/hwpml/2011/paragraph">'
                '<hp:p><hp:run><hp:t>E2E-V2</hp:t></hp:run></hp:p></hs:sec>',
            )
            archive.writestr(
                "Contents/content.hpf",
                '<opf:package xmlns:opf="http://www.idpf.org/2007/opf/">'
                '<opf:manifest><opf:item id="header" href="header.xml"/>'
                '<opf:item id="section0" href="section0.xml"/></opf:manifest>'
                '<opf:spine><opf:itemref idref="section0"/></opf:spine></opf:package>',
            )

    template = tmp_path / "template.hwpx"
    write_hwpx(template)
    monkeypatch.setenv("HWP_TEMPLATE_ROOT", str(tmp_path))
    monkeypatch.setenv("HWP_AGENT_TOKEN", "test-token")
    with sqlite_engine.begin() as conn:
        result = conn.execute(
            hwp_templates.insert().values(
                template_key="bid_form",
                kind="bid_form",
                name="Test",
                template_path=str(template),
                active=True,
                created_at=service.now(),
                updated_at=service.now(),
            )
        )
        conn.execute(
            document_field_mappings.insert().values(
                template_id=result.inserted_primary_key[0],
                hwp_field_name="notice",
                context_path="notice.notice_no",
                value_type="string",
                required=True,
                transform="none",
                sort_order=1,
                active=True,
                created_at=service.now(),
                updated_at=service.now(),
            )
        )
    assert client.post("/v2/notices/E2E-V2/sync", json={"version": 1}).status_code == 200

    class Worker:
        def __init__(self, **kwargs):
            assert kwargs["retries"] == 1

        def health(self):
            return True

        def inspect_fields(self, path):
            import hashlib
            from pathlib import Path

            return {"sha256": hashlib.sha256(Path(path).read_bytes()).hexdigest(),
                    "fields": ["notice"], "field_occurrences": [{"field": "notice", "occurrence": 0}]}

        def put_fields(self, **kwargs):
            from pathlib import Path

            assert kwargs["values"]["notice"] == "E2E-V2"
            write_hwpx(Path(kwargs["output_path"]))
            return PutFieldsOutcome(
                output_path=kwargs["output_path"],
                replaced=["notice"],
                missing=[],
                remaining_placeholders=["unmapped"],
                raw={},
            )

    monkeypatch.setattr("api.workspace.execution.HwpAgentClient", Worker)
    job = enqueue(client, kind="hwp", version=2).json()
    run_once(sqlite_engine)
    result = client.get(f"/v2/jobs/{job['id']}").json()
    assert result["status"] == "succeeded", result
    assert result["quality"] == "warning"
    assert result["result"]["remaining_placeholders"] == ["unmapped"]
    assert result["result"]["replaced"] == ["notice"]
    assert result["result"]["validation"]["structure"] == "passed"
    assert result["result"]["validation"]["format"] == "hwpx"
    assert result["result"]["verified"] is False
    assert len(result["result"]["sha256"]) == 64


def test_v2_auth_gate(client, workspace, monkeypatch):
    from api.config import settings

    monkeypatch.setattr(settings, "api_key", "local-test-key")
    assert client.get("/v2/notices").status_code == 403
    assert client.get("/v2/notices", headers={"X-API-Key": "local-test-key"}).status_code == 200


def test_list_notices_marks_workspaces_and_filters(client, workspace):
    """R2 regression: list query moved to api.workspace.queries keeps the response shape."""
    registered = client.post("/v2/notices", json={
        "notice_no": "E2E-LIST", "title": "목록 전용 공고", "source": "MANUAL_E2E", "raw": {},
    })
    assert registered.status_code == 200, registered.text
    assert client.post("/v2/notices", json={
        "notice_no": "E2E-LIST", "title": "목록 전용 공고", "source": "MANUAL_E2E", "raw": {},
    }).json() == {"notice_no": "E2E-LIST", "already_exists": True}

    body = client.get("/v2/notices").json()
    assert body["page"] == 1 and body["has_more"] is False
    by_no = {item["notice_no"]: item for item in body["items"]}
    assert by_no["E2E-V2"]["workspace"]["notice_no"] == "E2E-V2"
    assert by_no["E2E-LIST"]["workspace"] is None
    assert by_no["E2E-LIST"]["title"] == "목록 전용 공고"

    filtered = client.get("/v2/notices", params={"q": "목록"}).json()
    assert [item["notice_no"] for item in filtered["items"]] == ["E2E-LIST"]
